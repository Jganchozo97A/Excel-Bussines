import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.workbook.defined_name import DefinedName
import datetime
import os

NAVY = "1F3B57"
TEAL = "2CA58D"
GOLD = "C9A24B"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D9DEE3"
WHITE = "FFFFFF"
DARK_TEXT = "1A1A1A"
RED_NEG = "B23B3B"
INPUT_FILL = "FFF6D9"
LOCKED_TEXT = "555555"
FONT_NAME = "Arial"

f_h1 = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
f_h2 = Font(name=FONT_NAME, size=12, bold=True, color=WHITE)
f_h3 = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
f_body = Font(name=FONT_NAME, size=10, color=DARK_TEXT)
f_body_b = Font(name=FONT_NAME, size=10, bold=True, color=DARK_TEXT)
f_kpi_label = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
f_kpi_value = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
f_note = Font(name=FONT_NAME, size=9, italic=True, color=LOCKED_TEXT)
f_input = Font(name=FONT_NAME, size=10, color="0000FF")

fill_header = PatternFill("solid", fgColor=NAVY)
fill_header2 = PatternFill("solid", fgColor=TEAL)
fill_light = PatternFill("solid", fgColor=LIGHT_GRAY)
fill_input = PatternFill("solid", fgColor=INPUT_FILL)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_subtotal = PatternFill("solid", fgColor="E4E9EF")

thin = Side(style="thin", color=MID_GRAY)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)


wb = Workbook()
wb.remove(wb.active)


def style_header_row(ws, row, col_start, col_end, fill=fill_header, font=f_h2, height=22):
    ws.row_dimensions[row].height = height
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align_center
        cell.border = border_all


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def lock_all(ws):
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def unlock_range(ws, cell_range):
    for row in ws[cell_range]:
        for cell in row:
            cell.protection = Protection(locked=False)


def build_portada(title, subtitle, features, price):
    ws = wb.create_sheet("Portada")
    set_col_widths(ws, [4, 22, 22, 22, 22, 22, 4])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    ws["B2"] = "LÍNEA DE PRODUCTOS DIGITALES · PLANTILLAS EXCEL"
    ws["B2"].font = Font(name=FONT_NAME, size=10, bold=True, color=TEAL)
    ws.merge_cells("B4:F6")
    ws["B4"] = title
    ws["B4"].font = Font(name=FONT_NAME, size=26, bold=True, color=NAVY)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells("B7:F8")
    ws["B7"] = subtitle
    ws["B7"].font = Font(name=FONT_NAME, size=11, color=DARK_TEXT)
    ws["B7"].alignment = align_wrap
    for r in range(2, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[4].height = 30
    row0 = 10
    ws.merge_cells(f"B{row0}:F{row0}")
    ws[f"B{row0}"] = "QUÉ INCLUYE ESTA PLANTILLA"
    ws[f"B{row0}"].font = f_h3
    for c in range(2, 7):
        ws.cell(row=row0, column=c).fill = fill_light
    ws.row_dimensions[row0].height = 20
    r = row0 + 1
    for feat in features:
        ws.merge_cells(f"B{r}:F{r}")
        ws[f"B{r}"] = f"✔  {feat}"
        ws[f"B{r}"].font = f_body
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = f"Producto digital · Precio orientativo: {price}  |  Licencia de uso comercial"
    ws[f"B{r}"].font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD)
    r += 2
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = "© Tu Marca de Plantillas · www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_instrucciones(steps):
    ws = wb.create_sheet("Instrucciones")
    set_col_widths(ws, [3, 46, 46, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:C2")
    ws["B2"] = "CÓMO USAR ESTA PLANTILLA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    r = 4
    for title, body in steps:
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = title
        ws[f"B{r}"].font = f_h3
        for cc in (2, 3):
            ws.cell(row=r, column=cc).fill = fill_light
        ws.row_dimensions[r].height = 18
        r += 1
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = body
        ws[f"B{r}"].font = f_body
        ws[f"B{r}"].alignment = align_wrap
        ws.row_dimensions[r].height = 42
        r += 2
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = "Soporte y otras plantillas de esta línea de productos: www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_configuracion(extra_rows=None):
    ws = wb.create_sheet("Configuracion")
    set_col_widths(ws, [3, 32, 20, 40, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:D2")
    ws["B2"] = "CONFIGURACIÓN Y MONEDA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    headers_row = 4
    ws.cell(row=headers_row, column=2, value="Parámetro")
    ws.cell(row=headers_row, column=3, value="Valor")
    ws.cell(row=headers_row, column=4, value="Notas")
    style_header_row(ws, headers_row, 2, 4)

    rows = [
        ("Nombre de moneda local", "Soles (PEN)", "Solo referencia. Escribe el nombre de tu moneda."),
        ("Símbolo de moneda local", "S/", "Símbolo que se usará en el Dashboard cuando la vista esté en LOCAL."),
        ("Tipo de cambio (1 USD =)", 3.75, "Cuántas unidades de tu moneda local equivalen a 1 USD."),
        ("Moneda de visualización", "Local", "Elige LOCAL o USD. Afecta Estado de Resultados y Dashboard."),
    ]
    if extra_rows:
        rows.extend(extra_rows)

    row_map = {}
    r = 5
    for label, value, note in rows:
        ws.cell(row=r, column=2, value=label).font = f_body_b
        c = ws.cell(row=r, column=3, value=value)
        c.font = f_input
        c.fill = fill_input
        c.alignment = align_center
        c.border = border_all
        ws.cell(row=r, column=4, value=note).font = f_note
        ws.cell(row=r, column=4).alignment = align_wrap
        ws.row_dimensions[r].height = 18
        ws.cell(row=r, column=2).border = border_all
        row_map[label] = r
        r += 1

    ROW_VIS = row_map["Moneda de visualización"]
    ROW_TC = row_map["Tipo de cambio (1 USD =)"]
    ROW_SIMBOLO = row_map["Símbolo de moneda local"]

    r_calc_label = r
    ws.cell(row=r_calc_label, column=2, value="Valores calculados (no editar)").font = f_h3
    for cc in (2, 3, 4):
        ws.cell(row=r_calc_label, column=cc).fill = fill_light
    ws.row_dimensions[r_calc_label].height = 18
    r += 1

    ROW_FACTOR = r
    ws.cell(row=ROW_FACTOR, column=2, value="Factor de conversión aplicado").font = f_body_b
    ws.cell(row=ROW_FACTOR, column=3, value=f'=IF(C{ROW_VIS}="USD",1/C{ROW_TC},1)')
    ws.cell(row=ROW_FACTOR, column=3).number_format = "0.0000"
    ws.cell(row=ROW_FACTOR, column=3).border = border_all
    ws.cell(row=ROW_FACTOR, column=4, value="Vale 1 si vista LOCAL; vale 1/Tipo de cambio si vista USD.").font = f_note
    ws.row_dimensions[ROW_FACTOR].height = 18
    r += 1

    ROW_SIMBOLO_ACTUAL = r
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=2, value="Símbolo de moneda actual").font = f_body_b
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3, value=f'=IF(C{ROW_VIS}="USD","US$",C{ROW_SIMBOLO})')
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).border = border_all
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).alignment = align_center
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=4, value="Símbolo que verás en el Dashboard según la moneda elegida.").font = f_note
    ws.row_dimensions[ROW_SIMBOLO_ACTUAL].height = 18
    r += 1

    note_row = r + 1
    ws.merge_cells(f"B{note_row}:D{note_row}")
    ws[f"B{note_row}"] = "Cambia 'Moneda de visualización' entre Local y USD para recalcular automáticamente toda la plantilla."
    ws[f"B{note_row}"].font = f_note
    ws[f"B{note_row}"].alignment = align_wrap

    lock_all(ws)
    for label in row_map:
        rr = row_map[label]
        unlock_range(ws, f"C{rr}:C{rr}")

    dv_moneda = DataValidation(type="list", formula1='"Local,USD"', allow_blank=False)
    ws.add_data_validation(dv_moneda)
    dv_moneda.add(ws[f"C{ROW_VIS}"])

    return dict(ROW_TC=ROW_TC, ROW_VIS=ROW_VIS, ROW_SIMBOLO=ROW_SIMBOLO,
                ROW_FACTOR=ROW_FACTOR, ROW_SIMBOLO_ACTUAL=ROW_SIMBOLO_ACTUAL, row_map=row_map)


MONTHS_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

TIPOS_CUENTA = ["Ingreso", "Costo de venta", "Gasto operativo", "Gasto administrativo",
                "Gasto financiero", "Impuesto"]

# =========================================================
build_portada(
    "CONTABILIDAD BÁSICA PARA EMPRENDEDORES",
    ("Lleva el registro contable de tu negocio sin complicaciones: clasifica cada movimiento con "
     "tu plan de cuentas y obtén tu Estado de Resultados (ventas, costos, gastos y utilidad) "
     "mes a mes — con panel de control, gráficos automáticos y selector de moneda (local / USD)."),
    [
        "Dashboard con indicadores clave (KPIs) y gráficos automáticos",
        "Fórmulas automáticas: nada que calcular a mano",
        "Estado de Resultados mensual y anual (ventas, costos, gastos, utilidad)",
        "Plan de cuentas editable con clasificación por tipo",
        "Gráficos: pastel (gastos por tipo), barras y evolución de la utilidad neta",
        "Listas desplegables y validación de datos en cada registro",
        "Selector de moneda: local o USD, con tipo de cambio editable",
        "Hoja de instrucciones paso a paso",
        "Celdas y fórmulas protegidas: solo editas donde debes",
        "Diseño profesional y ejemplos precargados",
    ],
    "$14.99 USD",
)

build_instrucciones([
    ("1. Configuración",
     "Define el símbolo de tu moneda local, el tipo de cambio, el año de análisis y si la "
     "plantilla se muestra en moneda LOCAL o USD."),
    ("2. Plan de Cuentas",
     "En 'Plan de Cuentas' revisa o edita las cuentas contables de tu negocio y su tipo (Ingreso, "
     "Costo de venta, Gasto operativo, Gasto administrativo, Gasto financiero o Impuesto)."),
    ("3. Transacciones",
     "En 'Transacciones' registra cada movimiento: fecha, cuenta (lista desplegable), descripción "
     "y monto. El tipo de cuenta se completa solo."),
    ("4. Estado de Resultados",
     "La hoja 'Estado de Resultados' calcula automáticamente, mes a mes, tus ventas, costo de "
     "venta, utilidad bruta, gastos operativos y administrativos, utilidad operativa, gastos "
     "financieros, impuestos y utilidad neta — como un estado de resultados real."),
    ("5. Dashboard",
     "La hoja 'Dashboard' muestra tus ventas, utilidad bruta, utilidad neta y margen neto del año, "
     "con gráficos que se actualizan solos."),
    ("6. Cambiar de moneda",
     "Cambia 'Moneda de visualización' en 'Configuración' entre Local y USD: todo se recalcula "
     "automáticamente."),
    ("7. Celdas protegidas",
     "Las hojas están protegidas para que no borres fórmulas por accidente. Solo puedes escribir en "
     "las celdas resaltadas en color crema / texto azul. Si necesitas editar otra celda, ve a Revisar → Desproteger hoja (no pide contraseña)."),
    ("Nota",
     "Esta plantilla es contabilidad básica de ingresos y gastos (Estado de Resultados), pensada "
     "para emprendedores y pequeños negocios. No reemplaza la contabilidad formal por partida "
     "doble ni el balance general; para eso consulta a un contador."),
])

cfg = build_configuracion(extra_rows=[
    ("Año de análisis", 2026, "Año que se usa como referencia en el Estado de Resultados."),
])
ROW_FACTOR = cfg["ROW_FACTOR"]
ROW_SIMBOLO_ACTUAL = cfg["ROW_SIMBOLO_ACTUAL"]
ROW_VIS = cfg["ROW_VIS"]
ROW_ANIO = cfg["row_map"]["Año de análisis"]

# =========================================================
# PLAN DE CUENTAS
# =========================================================
ws = wb.create_sheet("Plan de Cuentas")
set_col_widths(ws, [12, 30, 22, 3])
ws.sheet_view.showGridLines = False
ws.merge_cells("A2:C2")
ws["A2"] = "PLAN DE CUENTAS"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26

pc_header = 4
ws.cell(row=pc_header, column=1, value="Código")
ws.cell(row=pc_header, column=2, value="Cuenta")
ws.cell(row=pc_header, column=3, value="Tipo")
style_header_row(ws, pc_header, 1, 3, height=20)

PC_FIRST = pc_header + 1
example_accounts = [
    ("4001", "Ventas de productos", "Ingreso"),
    ("4002", "Ventas de servicios", "Ingreso"),
    ("5001", "Costo de mercadería vendida", "Costo de venta"),
    ("6001", "Sueldos y salarios", "Gasto operativo"),
    ("6002", "Alquiler de local", "Gasto operativo"),
    ("6003", "Servicios (luz, agua, internet)", "Gasto operativo"),
    ("6004", "Marketing y publicidad", "Gasto administrativo"),
    ("6005", "Útiles y suministros de oficina", "Gasto administrativo"),
    ("7001", "Intereses y comisiones bancarias", "Gasto financiero"),
    ("8001", "Impuesto a la renta", "Impuesto"),
]
PC_LAST = PC_FIRST + len(example_accounts) - 1

for i, (code, name, tipo) in enumerate(example_accounts):
    rr = PC_FIRST + i
    ws.cell(row=rr, column=1, value=code)
    ws.cell(row=rr, column=2, value=name)
    ws.cell(row=rr, column=3, value=tipo)
    for c in range(1, 4):
        cell = ws.cell(row=rr, column=c)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        cell.alignment = align_left if c != 1 else align_center

dv_tipo_cuenta = DataValidation(type="list", formula1='"' + ",".join(TIPOS_CUENTA) + '"', allow_blank=True)
ws.add_data_validation(dv_tipo_cuenta)
dv_tipo_cuenta.add(f"C{PC_FIRST}:C{PC_LAST}")

note_row = PC_LAST + 2
ws.merge_cells(f"A{note_row}:C{note_row}")
ws[f"A{note_row}"] = "Puedes renombrar cuentas o agregar filas; el tipo determina en qué línea del Estado de Resultados aparece."
ws[f"A{note_row}"].font = f_note
ws[f"A{note_row}"].alignment = align_wrap

lock_all(ws)
unlock_range(ws, f"A{PC_FIRST}:C{PC_LAST}")

PC_NAME_RANGE = f"'Plan de Cuentas'!$B${PC_FIRST}:$B${PC_LAST}"
PC_TIPO_RANGE = f"'Plan de Cuentas'!$C${PC_FIRST}:$C${PC_LAST}"

# =========================================================
# TRANSACCIONES
# =========================================================
ws = wb.create_sheet("Transacciones")
set_col_widths(ws, [14, 30, 30, 16, 20])
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"

ws.merge_cells("A2:E2")
ws["A2"] = "REGISTRO DE TRANSACCIONES"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("A3:E3")
ws["A3"] = "Completa las columnas A a D para cada movimiento. La columna E (tipo) se calcula automáticamente."
ws["A3"].font = f_note

T_HEADER = 4
t_headers = ["Fecha", "Cuenta", "Descripción", "Monto", "Tipo de cuenta"]
for i, h in enumerate(t_headers, start=1):
    ws.cell(row=T_HEADER, column=i, value=h)
style_header_row(ws, T_HEADER, 1, 5, height=20)

T_FIRST = T_HEADER + 1
T_LAST = T_FIRST + 199

example_tx = [
    (datetime.date(2026, 1, 5), "Ventas de productos", "Ventas de la semana", 4200),
    (datetime.date(2026, 1, 8), "Costo de mercadería vendida", "Costo de productos vendidos", 1800),
    (datetime.date(2026, 1, 10), "Sueldos y salarios", "Planilla de enero", 1200),
    (datetime.date(2026, 1, 12), "Alquiler de local", "Alquiler de enero", 800),
    (datetime.date(2026, 1, 15), "Marketing y publicidad", "Campaña en redes sociales", 250),
    (datetime.date(2026, 1, 20), "Servicios (luz, agua, internet)", "Recibos del mes", 180),
    (datetime.date(2026, 1, 28), "Intereses y comisiones bancarias", "Comisiones bancarias", 40),
    (datetime.date(2026, 2, 5), "Ventas de productos", "Ventas de la semana", 4500),
    (datetime.date(2026, 2, 6), "Ventas de servicios", "Servicio de instalación", 600),
    (datetime.date(2026, 2, 8), "Costo de mercadería vendida", "Costo de productos vendidos", 1900),
    (datetime.date(2026, 2, 10), "Sueldos y salarios", "Planilla de febrero", 1200),
    (datetime.date(2026, 2, 12), "Alquiler de local", "Alquiler de febrero", 800),
    (datetime.date(2026, 2, 18), "Útiles y suministros de oficina", "Papelería", 60),
    (datetime.date(2026, 2, 27), "Impuesto a la renta", "Pago a cuenta mensual", 300),
    (datetime.date(2026, 3, 5), "Ventas de productos", "Ventas de la semana", 4800),
    (datetime.date(2026, 3, 8), "Costo de mercadería vendida", "Costo de productos vendidos", 2000),
    (datetime.date(2026, 3, 10), "Sueldos y salarios", "Planilla de marzo", 1200),
    (datetime.date(2026, 3, 12), "Alquiler de local", "Alquiler de marzo", 800),
    (datetime.date(2026, 3, 15), "Marketing y publicidad", "Campaña en redes sociales", 300),
    (datetime.date(2026, 3, 27), "Impuesto a la renta", "Pago a cuenta mensual", 320),
]

for i in range(T_LAST - T_FIRST + 1):
    rr = T_FIRST + i
    if i < len(example_tx):
        fecha, cuenta, desc, monto = example_tx[i]
        ws.cell(row=rr, column=1, value=fecha).number_format = "DD/MM/YYYY"
        ws.cell(row=rr, column=2, value=cuenta)
        ws.cell(row=rr, column=3, value=desc)
        ws.cell(row=rr, column=4, value=monto)
    else:
        ws.cell(row=rr, column=1).number_format = "DD/MM/YYYY"

    for col in (1, 2, 3, 4):
        cell = ws.cell(row=rr, column=col)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        if col == 4:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
        elif col == 1:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

    tcell = ws.cell(row=rr, column=5,
                     value=f'=IF(B{rr}="","",IFERROR(INDEX({PC_TIPO_RANGE},MATCH(B{rr},{PC_NAME_RANGE},0)),"Cuenta no encontrada"))')
    tcell.font = f_body
    tcell.fill = fill_light
    tcell.border = border_all
    tcell.alignment = align_left

dv_cuenta = DataValidation(type="list", formula1=f"={PC_NAME_RANGE}", allow_blank=True)
ws.add_data_validation(dv_cuenta)
dv_cuenta.add(f"B{T_FIRST}:B{T_LAST}")

dv_monto = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                           allow_blank=True, showErrorMessage=True,
                           errorTitle="Monto inválido", error="Ingresa un número igual o mayor a 0.")
ws.add_data_validation(dv_monto)
dv_monto.add(f"D{T_FIRST}:D{T_LAST}")

lock_all(ws)
unlock_range(ws, f"A{T_FIRST}:D{T_LAST}")

TX_FECHA_RANGE = f"Transacciones!$A${T_FIRST}:$A${T_LAST}"
TX_CUENTA_RANGE = f"Transacciones!$B${T_FIRST}:$B${T_LAST}"
TX_MONTO_RANGE = f"Transacciones!$D${T_FIRST}:$D${T_LAST}"
TX_TIPO_RANGE = f"Transacciones!$E${T_FIRST}:$E${T_LAST}"

# =========================================================
# ESTADO DE RESULTADOS
# =========================================================
ws = wb.create_sheet("Estado de Resultados")
set_col_widths(ws, [32] + [11] * 12 + [13])
ws.sheet_view.showGridLines = False
ws.freeze_panes = "B6"

ws.merge_cells("A2:N2")
ws["A2"] = "ESTADO DE RESULTADOS"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("A3:N3")
ws["A3"] = '=" Moneda: " & Configuracion!$C$8 & "   |   Año: " & Configuracion!$C$' + str(ROW_ANIO)
ws["A3"].font = f_note

er_header = 5
ws.cell(row=er_header, column=1, value="Concepto")
for i, m in enumerate(MONTHS_ES):
    ws.cell(row=er_header, column=2 + i, value=m[:3])
ws.cell(row=er_header, column=14, value="Total anual")
style_header_row(ws, er_header, 1, 14, height=20)


def month_end_expr(mnum):
    return f'EOMONTH(DATE(Configuracion!$C${ROW_ANIO},{mnum},1),0)'


def month_start_expr(mnum):
    return f'DATE(Configuracion!$C${ROW_ANIO},{mnum},1)'


def sumifs_tipo(tipo_literal, mnum):
    return (f'SUMIFS({TX_MONTO_RANGE},{TX_TIPO_RANGE},"{tipo_literal}",'
            f'{TX_FECHA_RANGE},">="&{month_start_expr(mnum)},{TX_FECHA_RANGE},"<="&{month_end_expr(mnum)})'
            f'*Configuracion!$C${ROW_FACTOR}')


rows_plan = [
    ("VENTAS", "line", "Ingreso"),
    ("Costo de venta", "line", "Costo de venta"),
    ("UTILIDAD BRUTA", "subtotal", None),
    ("Gastos operativos", "line", "Gasto operativo"),
    ("Gastos administrativos", "line", "Gasto administrativo"),
    ("UTILIDAD OPERATIVA", "subtotal", None),
    ("Gastos financieros", "line", "Gasto financiero"),
    ("Impuestos", "line", "Impuesto"),
    ("UTILIDAD NETA", "subtotal", None),
]

r = er_header + 1
row_of = {}
for label, kind, tipo in rows_plan:
    row_of[label] = r
    is_subtotal = kind == "subtotal"
    label_cell = ws.cell(row=r, column=1, value=label)
    label_cell.font = f_body_b if is_subtotal else f_body
    label_cell.border = border_all
    label_cell.fill = fill_subtotal if is_subtotal else fill_white

    for m in range(12):
        col = 2 + m
        mnum = m + 1
        cell = ws.cell(row=r, column=col)
        cell.border = border_all
        cell.number_format = "#,##0"
        cell.alignment = align_center
        if is_subtotal:
            cell.font = f_body_b
            cell.fill = fill_subtotal
        else:
            cell.font = f_body
            if tipo == "Ingreso":
                cell.value = f"={sumifs_tipo('Ingreso', mnum)}"
            else:
                cell.value = f"={sumifs_tipo(tipo, mnum)}"

    total_cell = ws.cell(row=r, column=14, value=f"=SUM(B{r}:M{r})")
    total_cell.border = border_all
    total_cell.number_format = "#,##0"
    total_cell.alignment = align_center
    total_cell.font = f_body_b if is_subtotal else f_body_b
    total_cell.fill = fill_subtotal if is_subtotal else fill_light
    r += 1

# Now fill subtotal formulas referencing the line rows
ventas_r = row_of["VENTAS"]
costo_r = row_of["Costo de venta"]
ub_r = row_of["UTILIDAD BRUTA"]
gop_r = row_of["Gastos operativos"]
gad_r = row_of["Gastos administrativos"]
uo_r = row_of["UTILIDAD OPERATIVA"]
gf_r = row_of["Gastos financieros"]
imp_r = row_of["Impuestos"]
un_r = row_of["UTILIDAD NETA"]

for col in list(range(2, 14)) + [14]:
    letter = get_column_letter(col)
    ws.cell(row=ub_r, column=col, value=f"={letter}{ventas_r}-{letter}{costo_r}")
    ws.cell(row=uo_r, column=col, value=f"={letter}{ub_r}-{letter}{gop_r}-{letter}{gad_r}")
    ws.cell(row=un_r, column=col, value=f"={letter}{uo_r}-{letter}{gf_r}-{letter}{imp_r}")

# Margin rows
margin_row_bruto = un_r + 2
ws.cell(row=margin_row_bruto, column=1, value="Margen bruto (%)").font = f_body
ws.cell(row=margin_row_bruto, column=1).border = border_all
margin_row_neto = margin_row_bruto + 1
ws.cell(row=margin_row_neto, column=1, value="Margen neto (%)").font = f_body
ws.cell(row=margin_row_neto, column=1).border = border_all
for col in list(range(2, 14)) + [14]:
    letter = get_column_letter(col)
    c1 = ws.cell(row=margin_row_bruto, column=col,
                 value=f'=IF({letter}{ventas_r}=0,0,{letter}{ub_r}/{letter}{ventas_r})')
    c1.number_format = "0.0%"
    c1.alignment = align_center
    c1.border = border_all
    c1.font = f_body
    c2 = ws.cell(row=margin_row_neto, column=col,
                 value=f'=IF({letter}{ventas_r}=0,0,{letter}{un_r}/{letter}{ventas_r})')
    c2.number_format = "0.0%"
    c2.alignment = align_center
    c2.border = border_all
    c2.font = f_body

lock_all(ws)

ER_SHEET = "Estado de Resultados"

# =========================================================
# DASHBOARD
# =========================================================
ws = wb.create_sheet("Dashboard")
set_col_widths(ws, [3, 16, 16, 16, 16, 16, 16, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:H2")
ws["B2"] = "DASHBOARD · CONTABILIDAD BÁSICA"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 28
ws.merge_cells("B3:H3")
ws["B3"] = '=" Moneda: " & Configuracion!$C$8 & "   |   Año: " & Configuracion!$C$' + str(ROW_ANIO)
ws["B3"].font = f_note

kpi_row = 5
card_defs = [
    ("VENTAS DEL AÑO", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT('{ER_SHEET}'!N{ventas_r},\"#,##0.00\")", TEAL),
    ("UTILIDAD BRUTA", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT('{ER_SHEET}'!N{ub_r},\"#,##0.00\")", NAVY),
    ("UTILIDAD NETA", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT('{ER_SHEET}'!N{un_r},\"#,##0.00\")", GOLD),
    ("MARGEN NETO", f'=IF(\'{ER_SHEET}\'!N{ventas_r}=0,"0.0%",TEXT(\'{ER_SHEET}\'!N{un_r}/\'{ER_SHEET}\'!N{ventas_r},"0.0%"))', RED_NEG),
]
col_starts = [2, 4, 6, 8]
for (label, formula, color), col0 in zip(card_defs, col_starts):
    ws.merge_cells(start_row=kpi_row, start_column=col0, end_row=kpi_row, end_column=col0 + 1)
    lc = ws.cell(row=kpi_row, column=col0, value=label)
    lc.font = f_kpi_label
    lc.alignment = align_center
    for c in range(col0, col0 + 2):
        ws.cell(row=kpi_row, column=c).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=kpi_row + 1, start_column=col0, end_row=kpi_row + 2, end_column=col0 + 1)
    vc = ws.cell(row=kpi_row + 1, column=col0, value=formula)
    vc.font = f_kpi_value
    vc.alignment = align_center
    for rr in (kpi_row + 1, kpi_row + 2):
        for c in range(col0, col0 + 2):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor=color)

ws.row_dimensions[kpi_row].height = 18
ws.row_dimensions[kpi_row + 1].height = 22
ws.row_dimensions[kpi_row + 2].height = 22

chart_top = kpi_row + 4

# Helper block for chart source data (a compact table so charts don't need 12 wide columns)
helper_row = chart_top - 1
ws.cell(row=1, column=1, value=None)  # no-op, keep structure tidy

pie = PieChart()
pie.title = "Gastos por tipo (año)"
gasto_labels_row = un_r  # not used directly; build small helper table instead
# Build a small helper table on Dashboard (hidden columns far right) for chart categories
helper_col = 11  # column K onward, off to the side but still on Dashboard sheet
ws.cell(row=4, column=helper_col, value="Tipo")
ws.cell(row=4, column=helper_col + 1, value="Monto")
gasto_tipos = [("Costo de venta", costo_r), ("Gastos operativos", gop_r),
               ("Gastos administrativos", gad_r), ("Gastos financieros", gf_r),
               ("Impuestos", imp_r)]
for i, (label, rref) in enumerate(gasto_tipos):
    rr = 5 + i
    ws.cell(row=rr, column=helper_col, value=label)
    ws.cell(row=rr, column=helper_col + 1, value=f"='{ER_SHEET}'!N{rref}")
    ws.cell(row=rr, column=helper_col + 1).number_format = "#,##0.00"
helper_last = 5 + len(gasto_tipos) - 1
ws.column_dimensions[get_column_letter(helper_col)].hidden = True
ws.column_dimensions[get_column_letter(helper_col + 1)].hidden = True

data_pie = Reference(ws, min_col=helper_col + 1, min_row=5, max_row=helper_last)
cats_pie = Reference(ws, min_col=helper_col, min_row=5, max_row=helper_last)
pie.add_data(data_pie, titles_from_data=False)
pie.set_categories(cats_pie)
pie.height = 9
pie.width = 15
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showVal = False
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showLegendKey = False
pie.dataLabels.showBubbleSize = False
ws.add_chart(pie, f"B{chart_top}")

bar = BarChart()
bar.type = "col"
bar.title = "Ventas vs. utilidad neta por mes"
bar.style = 10
mcats = Reference(wb[ER_SHEET], min_col=2, max_col=13, min_row=er_header, max_row=er_header)

bdata_ventas = Reference(wb[ER_SHEET], min_col=2, max_col=13, min_row=ventas_r, max_row=ventas_r)
bar.add_data(bdata_ventas, titles_from_data=False, from_rows=True)
bar.series[-1].tx = SeriesLabel(v="Ventas")

bdata_utilidad = Reference(wb[ER_SHEET], min_col=2, max_col=13, min_row=un_r, max_row=un_r)
bar.add_data(bdata_utilidad, titles_from_data=False, from_rows=True)
bar.series[-1].tx = SeriesLabel(v="Utilidad neta")

bar.set_categories(mcats)
bar.height = 9
bar.width = 15
ws.add_chart(bar, f"E{chart_top}")

line = LineChart()
line.title = "Evolución de la utilidad neta"
ldata = Reference(wb[ER_SHEET], min_col=2, max_col=13, min_row=un_r, max_row=un_r)
line.add_data(ldata, titles_from_data=False, from_rows=True)
line.series[-1].tx = SeriesLabel(v="Utilidad neta")
line.set_categories(mcats)
line.height = 9
line.width = 30
ws.add_chart(line, f"B{chart_top + 19}")

lock_all(ws)

# =========================================================
wb.defined_names["TipoCambio"] = DefinedName("TipoCambio", attr_text=f"Configuracion!$C${cfg['ROW_TC']}")
wb.defined_names["MonedaVisualizacion"] = DefinedName("MonedaVisualizacion", attr_text=f"Configuracion!$C${ROW_VIS}")
wb.defined_names["FactorConversion"] = DefinedName("FactorConversion", attr_text=f"Configuracion!$C${ROW_FACTOR}")

order = ["Portada", "Instrucciones", "Configuracion", "Plan de Cuentas", "Transacciones",
          "Estado de Resultados", "Dashboard"]
wb._sheets = [wb[name] for name in order]
tab_colors = {
    "Portada": NAVY, "Instrucciones": TEAL, "Configuracion": GOLD, "Plan de Cuentas": GOLD,
    "Transacciones": TEAL, "Estado de Resultados": NAVY, "Dashboard": NAVY,
}
for name, color in tab_colors.items():
    wb[name].sheet_properties.tabColor = color

wb.active = wb.index(wb["Dashboard"])
wb.calculation.fullCalcOnLoad = True

for name in order:
    psheet = wb[name]
    psheet.page_setup.orientation = "landscape"
    psheet.page_setup.fitToWidth = 1
    psheet.page_setup.fitToHeight = 0
    psheet.sheet_properties.pageSetUpPr.fitToPage = True
    psheet.page_margins.left = 0.4
    psheet.page_margins.right = 0.4
    psheet.page_margins.top = 0.5
    psheet.page_margins.bottom = 0.5

out_path = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "productos", "06-contabilidad-basica", "Contabilidad-Basica-para-Emprendedores.xlsx",
)
os.makedirs(os.path.dirname(out_path), exist_ok=True)
wb.save(out_path)
print("Saved:", out_path)
print("Key rows:", dict(ventas_r=ventas_r, ub_r=ub_r, uo_r=uo_r, un_r=un_r))
