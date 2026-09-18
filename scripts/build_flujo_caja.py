import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.defined_name import DefinedName
import os

NAVY = "1F3B57"
TEAL = "2CA58D"
GOLD = "C9A24B"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D9DEE3"
WHITE = "FFFFFF"
DARK_TEXT = "1A1A1A"
RED_NEG = "B23B3B"
INPUT_FILL = "FFF6D9"
LOCKED_TEXT = "555555"
FONT_NAME = "Arial"

f_h1 = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
f_h2 = Font(name=FONT_NAME, size=12, bold=True, color=WHITE)
f_h3 = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
f_body = Font(name=FONT_NAME, size=10, color=DARK_TEXT)
f_body_b = Font(name=FONT_NAME, size=10, bold=True, color=DARK_TEXT)
f_kpi_label = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
f_kpi_value = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
f_note = Font(name=FONT_NAME, size=9, italic=True, color=LOCKED_TEXT)
f_input = Font(name=FONT_NAME, size=10, color="0000FF")

fill_header = PatternFill("solid", fgColor=NAVY)
fill_header2 = PatternFill("solid", fgColor=TEAL)
fill_light = PatternFill("solid", fgColor=LIGHT_GRAY)
fill_input = PatternFill("solid", fgColor=INPUT_FILL)
fill_white = PatternFill("solid", fgColor=WHITE)

thin = Side(style="thin", color=MID_GRAY)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

PROTECT_PASSWORD = "plantilla2026"

wb = Workbook()
wb.remove(wb.active)


def style_header_row(ws, row, col_start, col_end, fill=fill_header, font=f_h2, height=22):
    ws.row_dimensions[row].height = height
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align_center
        cell.border = border_all


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def lock_all(ws):
    ws.protection.sheet = True
    ws.protection.password = PROTECT_PASSWORD
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def unlock_range(ws, cell_range):
    for row in ws[cell_range]:
        for cell in row:
            cell.protection = Protection(locked=False)


def build_portada(title, subtitle, features, price):
    ws = wb.create_sheet("Portada")
    set_col_widths(ws, [4, 22, 22, 22, 22, 22, 4])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    ws["B2"] = "LÍNEA DE PRODUCTOS DIGITALES · PLANTILLAS EXCEL"
    ws["B2"].font = Font(name=FONT_NAME, size=10, bold=True, color=TEAL)
    ws.merge_cells("B4:F6")
    ws["B4"] = title
    ws["B4"].font = Font(name=FONT_NAME, size=28, bold=True, color=NAVY)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells("B7:F8")
    ws["B7"] = subtitle
    ws["B7"].font = Font(name=FONT_NAME, size=11, color=DARK_TEXT)
    ws["B7"].alignment = align_wrap
    for r in range(2, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[4].height = 30
    row0 = 10
    ws.merge_cells(f"B{row0}:F{row0}")
    ws[f"B{row0}"] = "QUÉ INCLUYE ESTA PLANTILLA"
    ws[f"B{row0}"].font = f_h3
    for c in range(2, 7):
        ws.cell(row=row0, column=c).fill = fill_light
    ws.row_dimensions[row0].height = 20
    r = row0 + 1
    for feat in features:
        ws.merge_cells(f"B{r}:F{r}")
        ws[f"B{r}"] = f"✔  {feat}"
        ws[f"B{r}"].font = f_body
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = f"Producto digital · Precio orientativo: {price}  |  Licencia de uso comercial"
    ws[f"B{r}"].font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD)
    r += 2
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = "© Tu Marca de Plantillas · www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_instrucciones(steps):
    ws = wb.create_sheet("Instrucciones")
    set_col_widths(ws, [3, 46, 46, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:C2")
    ws["B2"] = "CÓMO USAR ESTA PLANTILLA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    r = 4
    for title, body in steps:
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = title
        ws[f"B{r}"].font = f_h3
        for cc in (2, 3):
            ws.cell(row=r, column=cc).fill = fill_light
        ws.row_dimensions[r].height = 18
        r += 1
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = body
        ws[f"B{r}"].font = f_body
        ws[f"B{r}"].alignment = align_wrap
        ws.row_dimensions[r].height = 42
        r += 2
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = "Soporte y otras plantillas de esta línea de productos: www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_configuracion(extra_rows=None):
    ws = wb.create_sheet("Configuracion")
    set_col_widths(ws, [3, 32, 20, 40, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:D2")
    ws["B2"] = "CONFIGURACIÓN Y MONEDA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    headers_row = 4
    ws.cell(row=headers_row, column=2, value="Parámetro")
    ws.cell(row=headers_row, column=3, value="Valor")
    ws.cell(row=headers_row, column=4, value="Notas")
    style_header_row(ws, headers_row, 2, 4)

    rows = [
        ("Nombre de moneda local", "Soles (PEN)", "Solo referencia. Escribe el nombre de tu moneda."),
        ("Símbolo de moneda local", "S/", "Símbolo que se usará en el Dashboard cuando la vista esté en LOCAL."),
        ("Tipo de cambio (1 USD =)", 3.75, "Cuántas unidades de tu moneda local equivalen a 1 USD."),
        ("Moneda de visualización", "Local", "Elige LOCAL o USD. Afecta Resumen y Dashboard."),
    ]
    if extra_rows:
        rows.extend(extra_rows)

    row_map = {}
    r = 5
    for label, value, note in rows:
        ws.cell(row=r, column=2, value=label).font = f_body_b
        c = ws.cell(row=r, column=3, value=value)
        c.font = f_input
        c.fill = fill_input
        c.alignment = align_center
        c.border = border_all
        ws.cell(row=r, column=4, value=note).font = f_note
        ws.cell(row=r, column=4).alignment = align_wrap
        ws.row_dimensions[r].height = 18
        ws.cell(row=r, column=2).border = border_all
        row_map[label] = r
        r += 1

    ROW_VIS = row_map["Moneda de visualización"]
    ROW_TC = row_map["Tipo de cambio (1 USD =)"]
    ROW_SIMBOLO = row_map["Símbolo de moneda local"]

    r_calc_label = r
    ws.cell(row=r_calc_label, column=2, value="Valores calculados (no editar)").font = f_h3
    for cc in (2, 3, 4):
        ws.cell(row=r_calc_label, column=cc).fill = fill_light
    ws.row_dimensions[r_calc_label].height = 18
    r += 1

    ROW_FACTOR = r
    ws.cell(row=ROW_FACTOR, column=2, value="Factor de conversión aplicado").font = f_body_b
    ws.cell(row=ROW_FACTOR, column=3, value=f'=IF(C{ROW_VIS}="USD",1/C{ROW_TC},1)')
    ws.cell(row=ROW_FACTOR, column=3).number_format = "0.0000"
    ws.cell(row=ROW_FACTOR, column=3).border = border_all
    ws.cell(row=ROW_FACTOR, column=4, value="Vale 1 si vista LOCAL; vale 1/Tipo de cambio si vista USD.").font = f_note
    ws.row_dimensions[ROW_FACTOR].height = 18
    r += 1

    ROW_SIMBOLO_ACTUAL = r
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=2, value="Símbolo de moneda actual").font = f_body_b
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3, value=f'=IF(C{ROW_VIS}="USD","US$",C{ROW_SIMBOLO})')
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).border = border_all
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).alignment = align_center
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=4, value="Símbolo que verás en el Dashboard según la moneda elegida.").font = f_note
    ws.row_dimensions[ROW_SIMBOLO_ACTUAL].height = 18
    r += 1

    note_row = r + 1
    ws.merge_cells(f"B{note_row}:D{note_row}")
    ws[f"B{note_row}"] = "Cambia 'Moneda de visualización' entre Local y USD para recalcular automáticamente toda la plantilla."
    ws[f"B{note_row}"].font = f_note
    ws[f"B{note_row}"].alignment = align_wrap

    lock_all(ws)
    for label in row_map:
        rr = row_map[label]
        unlock_range(ws, f"C{rr}:C{rr}")

    dv_moneda = DataValidation(type="list", formula1='"Local,USD"', allow_blank=False)
    ws.add_data_validation(dv_moneda)
    dv_moneda.add(ws[f"C{ROW_VIS}"])

    return dict(ROW_TC=ROW_TC, ROW_VIS=ROW_VIS, ROW_SIMBOLO=ROW_SIMBOLO,
                ROW_FACTOR=ROW_FACTOR, ROW_SIMBOLO_ACTUAL=ROW_SIMBOLO_ACTUAL, row_map=row_map)


MONTHS_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
MONTHS_SHORT = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

categorias = [
    "Ventas al contado", "Cobros de cuentas por cobrar", "Otros ingresos",
    "Compra de mercadería/insumos", "Sueldos y planilla", "Alquiler",
    "Servicios (luz, agua, internet)", "Marketing y publicidad", "Impuestos",
    "Mantenimiento y otros gastos", "Otros egresos",
]
N_INCOME = 3
N_CATS = len(categorias)

# =========================================================
build_portada(
    "FLUJO DE CAJA PARA NEGOCIOS",
    ("Proyecta y registra las entradas y salidas de efectivo de tu negocio mes a mes, y visualiza "
     "tu saldo de caja en todo momento — con panel de control, gráficos automáticos y selector de "
     "moneda (local / USD)."),
    [
        "Dashboard con indicadores clave (KPIs) y gráficos automáticos",
        "Fórmulas automáticas: nada que calcular a mano",
        "Gráficos: pastel (egresos por categoría), barras y evolución del saldo de caja",
        "Resumen dinámico tipo tabla dinámica por categoría y por mes",
        "Saldo de caja proyectado mes a mes (arrastre automático)",
        "Listas desplegables y validación de datos en cada registro",
        "Selector de moneda: local o USD, con tipo de cambio editable",
        "Hoja de instrucciones paso a paso",
        "Celdas y fórmulas protegidas: solo editas donde debes",
        "Diseño profesional y ejemplos precargados",
    ],
    "$12.99 USD",
)

build_instrucciones([
    ("1. Configuración",
     "Define el símbolo de tu moneda local, el tipo de cambio, el saldo de caja inicial de tu "
     "negocio y si la plantilla se muestra en moneda LOCAL o USD."),
    ("2. Categorías",
     "En la hoja 'Categorías' revisa o edita las categorías de ingreso/egreso de efectivo de tu "
     "negocio."),
    ("3. Flujo de Caja",
     "En la hoja 'Flujo de Caja' registra, mes a mes, cuánto efectivo entró y salió en cada "
     "categoría. La columna 'Total anual' se calcula sola."),
    ("4. Resumen automático",
     "La hoja 'Resumen' calcula tus ingresos y egresos totales por mes, tu flujo neto y tu saldo "
     "de caja proyectado (arrastrando el saldo inicial mes a mes), como una tabla dinámica."),
    ("5. Dashboard",
     "La hoja 'Dashboard' muestra tu saldo de caja actual, tu flujo neto y gráficos que se "
     "actualizan solos: ingresos vs. egresos por mes, egresos por categoría y la curva de tu "
     "saldo de caja."),
    ("6. Cambiar de moneda",
     "Cambia 'Moneda de visualización' en 'Configuración' entre Local y USD: todo se recalcula "
     "automáticamente."),
    ("7. Celdas protegidas",
     "Las hojas están protegidas para que no borres fórmulas por accidente. Solo puedes escribir en "
     "las celdas resaltadas en color crema / texto azul. Contraseña de desprotección: plantilla2026."),
])

cfg = build_configuracion(extra_rows=[
    ("Saldo de caja inicial", 5000, "Efectivo disponible al comenzar el año (caja + bancos)."),
    ("Año de análisis", 2026, "Año que se usa como referencia en Flujo de Caja."),
])
ROW_FACTOR = cfg["ROW_FACTOR"]
ROW_SIMBOLO_ACTUAL = cfg["ROW_SIMBOLO_ACTUAL"]
ROW_VIS = cfg["ROW_VIS"]
ROW_SALDO_INICIAL = cfg["row_map"]["Saldo de caja inicial"]
ROW_ANIO = cfg["row_map"]["Año de análisis"]

# =========================================================
# CATEGORIAS
# =========================================================
ws = wb.create_sheet("Categorias")
set_col_widths(ws, [3, 34, 3])
ws.sheet_view.showGridLines = False
ws.merge_cells("B2:B2")
ws["B2"] = "CATEGORÍAS DE FLUJO DE CAJA"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 26

ws.cell(row=4, column=2, value="Categoría (ingreso/egreso)")
style_header_row(ws, 4, 2, 2)

CAT_FIRST_ROW = 5
CAT_LAST_ROW = CAT_FIRST_ROW + len(categorias) - 1
for i, cat in enumerate(categorias):
    rr = CAT_FIRST_ROW + i
    c = ws.cell(row=rr, column=2, value=cat)
    c.font = f_input
    c.fill = fill_input
    c.border = border_all
    ws.row_dimensions[rr].height = 16

note_row = CAT_LAST_ROW + 2
ws.merge_cells(f"B{note_row}:B{note_row}")
ws[f"B{note_row}"] = "Estas categorías deben coincidir, en el mismo orden, con las filas de 'Flujo de Caja'."
ws[f"B{note_row}"].font = f_note
ws[f"B{note_row}"].alignment = align_wrap

lock_all(ws)
unlock_range(ws, f"B{CAT_FIRST_ROW}:B{CAT_LAST_ROW}")

# =========================================================
# FLUJO DE CAJA (grid categoría x mes)
# =========================================================
ws = wb.create_sheet("Flujo de Caja")
set_col_widths(ws, [32] + [10] * 12 + [13])
ws.sheet_view.showGridLines = False
ws.freeze_panes = "B5"

ws.merge_cells("A2:N2")
ws["A2"] = "FLUJO DE CAJA MENSUAL (por categoría)"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26

header_row = 4
ws.cell(row=header_row, column=1, value="Categoría")
for i, m in enumerate(MONTHS_SHORT):
    ws.cell(row=header_row, column=2 + i, value=m)
ws.cell(row=header_row, column=14, value="Total anual")
style_header_row(ws, header_row, 1, 14, height=20)

FC_FIRST = header_row + 1
FC_LAST = FC_FIRST + N_CATS - 1

example_income = {
    "Ventas al contado": [8000, 8200, 8500, 8300, 8600, 8700, 8900, 8800, 9000, 9200, 9500, 10500],
    "Cobros de cuentas por cobrar": [1500, 1600, 1400, 1550, 1600, 1650, 1700, 1650, 1800, 1750, 1900, 2000],
    "Otros ingresos": [200, 0, 300, 0, 150, 0, 0, 400, 0, 0, 250, 0],
}
example_expense = {
    "Compra de mercadería/insumos": [3200, 3300, 3400, 3350, 3450, 3500, 3600, 3550, 3650, 3700, 3800, 4200],
    "Sueldos y planilla": [3500, 3500, 3500, 3500, 3500, 3500, 3500, 3500, 3500, 3500, 3500, 4200],
    "Alquiler": [1200, 1200, 1200, 1200, 1200, 1200, 1200, 1200, 1200, 1200, 1200, 1200],
    "Servicios (luz, agua, internet)": [280, 290, 270, 300, 310, 320, 330, 310, 300, 290, 300, 320],
    "Marketing y publicidad": [400, 400, 500, 400, 450, 400, 500, 450, 400, 500, 600, 800],
    "Impuestos": [600, 0, 0, 600, 0, 0, 600, 0, 0, 600, 0, 0],
    "Mantenimiento y otros gastos": [150, 100, 200, 150, 100, 180, 150, 120, 150, 200, 180, 250],
    "Otros egresos": [100, 80, 120, 90, 100, 110, 90, 100, 120, 100, 110, 150],
}

for i, cat in enumerate(categorias):
    rr = FC_FIRST + i
    is_income = i < N_INCOME
    ws.cell(row=rr, column=1, value=f"=Categorias!$B${CAT_FIRST_ROW + i}")
    ws.cell(row=rr, column=1).font = f_body_b if is_income else f_body
    ws.cell(row=rr, column=1).border = border_all
    ws.cell(row=rr, column=1).fill = fill_light if is_income else fill_white

    example_row = example_income if is_income else example_expense
    example_vals = example_row.get(cat, [None] * 12)

    for m in range(12):
        cell = ws.cell(row=rr, column=2 + m)
        val = example_vals[m] if m < len(example_vals) else None
        if val is not None:
            cell.value = val
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        cell.number_format = "#,##0"
        cell.alignment = align_center

    tot = ws.cell(row=rr, column=14, value=f"=SUM(B{rr}:M{rr})")
    tot.number_format = "#,##0"
    tot.font = f_body_b
    tot.fill = fill_light
    tot.border = border_all
    tot.alignment = align_center

total_row = FC_LAST + 1
ws.cell(row=total_row, column=1, value="TOTAL").font = f_body_b
ws.cell(row=total_row, column=1).border = border_all
ws.cell(row=total_row, column=1).fill = fill_light
for m in range(12):
    col = 2 + m
    letter = get_column_letter(col)
    c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{FC_FIRST}:{letter}{FC_LAST})")
    c.number_format = "#,##0"
    c.font = f_body_b
    c.fill = fill_light
    c.border = border_all
    c.alignment = align_center
c = ws.cell(row=total_row, column=14, value=f"=SUM(N{FC_FIRST}:N{FC_LAST})")
c.number_format = "#,##0"
c.font = f_body_b
c.fill = fill_light
c.border = border_all
c.alignment = align_center

dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                         allow_blank=True, showErrorMessage=True,
                         errorTitle="Monto inválido", error="Ingresa un número igual o mayor a 0.")
ws.add_data_validation(dv_num)
dv_num.add(f"B{FC_FIRST}:M{FC_LAST}")

lock_all(ws)
unlock_range(ws, f"B{FC_FIRST}:M{FC_LAST}")

FC_INCOME_FIRST = FC_FIRST
FC_INCOME_LAST = FC_FIRST + N_INCOME - 1
FC_EXPENSE_FIRST = FC_INCOME_LAST + 1
FC_EXPENSE_LAST = FC_LAST

# =========================================================
# RESUMEN
# =========================================================
ws = wb.create_sheet("Resumen")
set_col_widths(ws, [3, 30, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:C2")
ws["B2"] = "RESUMEN DINÁMICO (se actualiza solo)"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("B3:C3")
ws["B3"] = '=" Moneda: " & Configuracion!$C$8'
ws["B3"].font = f_note

r0 = 5
ws.cell(row=r0, column=2, value="TOTAL ANUAL POR CATEGORÍA")
ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=3)
ws.cell(row=r0, column=2).font = f_h2
for c in range(2, 4):
    ws.cell(row=r0, column=c).fill = fill_header2
ws.row_dimensions[r0].height = 20

hr = r0 + 1
ws.cell(row=hr, column=2, value="Categoría")
ws.cell(row=hr, column=3, value="Total anual")
style_header_row(ws, hr, 2, 3, height=18)

CAT_SUMMARY_FIRST = hr + 1
for i, cat in enumerate(categorias):
    rr = CAT_SUMMARY_FIRST + i
    fc_row = FC_FIRST + i
    ws.cell(row=rr, column=2, value=f"=Categorias!$B${CAT_FIRST_ROW + i}")
    ws.cell(row=rr, column=3, value=f"='Flujo de Caja'!$N${fc_row}*Configuracion!$C${ROW_FACTOR}")
    for c in (2, 3):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c == 3:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
CAT_SUMMARY_LAST = CAT_SUMMARY_FIRST + N_CATS - 1
CAT_EXPENSE_FIRST = CAT_SUMMARY_FIRST + N_INCOME

# ---- Monthly cash flow with running balance ----
r1 = CAT_SUMMARY_LAST + 3
ws.cell(row=r1, column=2, value="FLUJO DE CAJA MENSUAL Y SALDO PROYECTADO")
ws.merge_cells(start_row=r1, start_column=2, end_row=r1, end_column=7)
ws.cell(row=r1, column=2).font = f_h2
for c in range(2, 8):
    ws.cell(row=r1, column=c).fill = fill_header2
ws.row_dimensions[r1].height = 20

hr2 = r1 + 1
mon_headers = ["Mes", "Ingresos", "Egresos", "Flujo neto", "Saldo inicial", "Saldo final"]
for i, h in enumerate(mon_headers):
    ws.cell(row=hr2, column=2 + i, value=h)
style_header_row(ws, hr2, 2, 7, height=18)

MONTH_FIRST = hr2 + 1
for i, mname in enumerate(MONTHS_ES):
    rr = MONTH_FIRST + i
    col_letter = get_column_letter(2 + i)
    ws.cell(row=rr, column=2, value=mname)
    ws.cell(row=rr, column=3,
            value=(f"=SUM('Flujo de Caja'!{col_letter}${FC_INCOME_FIRST}:{col_letter}${FC_INCOME_LAST})"
                   f"*Configuracion!$C${ROW_FACTOR}"))
    ws.cell(row=rr, column=4,
            value=(f"=SUM('Flujo de Caja'!{col_letter}${FC_EXPENSE_FIRST}:{col_letter}${FC_EXPENSE_LAST})"
                   f"*Configuracion!$C${ROW_FACTOR}"))
    ws.cell(row=rr, column=5, value=f"=C{rr}-D{rr}")
    if i == 0:
        ws.cell(row=rr, column=6, value=f"=Configuracion!$C${ROW_SALDO_INICIAL}*Configuracion!$C${ROW_FACTOR}")
    else:
        ws.cell(row=rr, column=6, value=f"=G{rr-1}")
    ws.cell(row=rr, column=7, value=f"=F{rr}+E{rr}")
    for c in range(2, 8):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c in (3, 4, 5, 6, 7):
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
MONTH_LAST = MONTH_FIRST + len(MONTHS_ES) - 1

tot_row = MONTH_LAST + 1
ws.cell(row=tot_row, column=2, value="TOTAL / SALDO FINAL DEL AÑO").font = f_body_b
for c, col_letter in zip((3, 4, 5), ("C", "D", "E")):
    cell = ws.cell(row=tot_row, column=c, value=f"=SUM({col_letter}{MONTH_FIRST}:{col_letter}{MONTH_LAST})")
    cell.number_format = "#,##0.00"
    cell.font = f_body_b
    cell.fill = fill_light
    cell.alignment = align_center
cell = ws.cell(row=tot_row, column=7, value=f"=G{MONTH_LAST}")
cell.number_format = "#,##0.00"
cell.font = f_body_b
cell.fill = fill_light
cell.alignment = align_center
for c in range(2, 8):
    ws.cell(row=tot_row, column=c).border = border_all
    ws.cell(row=tot_row, column=c).fill = fill_light
RES_TOTAL_ROW = tot_row

lock_all(ws)

# =========================================================
# DASHBOARD
# =========================================================
ws = wb.create_sheet("Dashboard")
set_col_widths(ws, [3, 16, 16, 16, 16, 16, 16, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:H2")
ws["B2"] = "DASHBOARD · FLUJO DE CAJA PARA NEGOCIOS"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 28
ws.merge_cells("B3:H3")
ws["B3"] = ('=" Moneda: " & Configuracion!$C$8 & "   |   Año: " & Configuracion!$C$'
            + str(ROW_ANIO))
ws["B3"].font = f_note

kpi_row = 5
card_defs = [
    ("SALDO DE CAJA INICIAL", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!F{MONTH_FIRST},\"#,##0.00\")", NAVY),
    ("TOTAL INGRESOS (AÑO)", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!C{RES_TOTAL_ROW},\"#,##0.00\")", TEAL),
    ("TOTAL EGRESOS (AÑO)", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!D{RES_TOTAL_ROW},\"#,##0.00\")", RED_NEG),
    ("SALDO DE CAJA FINAL", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!G{RES_TOTAL_ROW},\"#,##0.00\")", GOLD),
]
col_starts = [2, 4, 6, 8]
for (label, formula, color), col0 in zip(card_defs, col_starts):
    ws.merge_cells(start_row=kpi_row, start_column=col0, end_row=kpi_row, end_column=col0 + 1)
    lc = ws.cell(row=kpi_row, column=col0, value=label)
    lc.font = f_kpi_label
    lc.alignment = align_center
    for c in range(col0, col0 + 2):
        ws.cell(row=kpi_row, column=c).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=kpi_row + 1, start_column=col0, end_row=kpi_row + 2, end_column=col0 + 1)
    vc = ws.cell(row=kpi_row + 1, column=col0, value=formula)
    vc.font = f_kpi_value
    vc.alignment = align_center
    for rr in (kpi_row + 1, kpi_row + 2):
        for c in range(col0, col0 + 2):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor=color)

ws.row_dimensions[kpi_row].height = 18
ws.row_dimensions[kpi_row + 1].height = 22
ws.row_dimensions[kpi_row + 2].height = 22

chart_top = kpi_row + 4

pie = PieChart()
pie.title = "Egresos por categoría (año)"
data_pie = Reference(wb["Resumen"], min_col=3, min_row=hr, max_row=CAT_SUMMARY_LAST)
cats_pie = Reference(wb["Resumen"], min_col=2, min_row=CAT_EXPENSE_FIRST, max_row=CAT_SUMMARY_LAST)
data_pie_expense = Reference(wb["Resumen"], min_col=3, min_row=CAT_EXPENSE_FIRST, max_row=CAT_SUMMARY_LAST)
pie.add_data(data_pie_expense, titles_from_data=False)
pie.set_categories(cats_pie)
pie.height = 9
pie.width = 15
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showVal = False
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showLegendKey = False
pie.dataLabels.showBubbleSize = False
ws.add_chart(pie, f"B{chart_top}")

bar = BarChart()
bar.type = "col"
bar.title = "Ingresos vs. egresos por mes"
bar.style = 10
bdata = Reference(wb["Resumen"], min_col=3, max_col=4, min_row=hr2, max_row=MONTH_LAST)
bcats = Reference(wb["Resumen"], min_col=2, min_row=MONTH_FIRST, max_row=MONTH_LAST)
bar.add_data(bdata, titles_from_data=True)
bar.set_categories(bcats)
bar.height = 9
bar.width = 15
ws.add_chart(bar, f"E{chart_top}")

line = LineChart()
line.title = "Evolución del saldo de caja"
ldata = Reference(wb["Resumen"], min_col=7, min_row=hr2, max_row=MONTH_LAST)
lcats = Reference(wb["Resumen"], min_col=2, min_row=MONTH_FIRST, max_row=MONTH_LAST)
line.add_data(ldata, titles_from_data=True)
line.set_categories(lcats)
line.height = 9
line.width = 30
ws.add_chart(line, f"B{chart_top + 19}")

lock_all(ws)

# =========================================================
wb.defined_names["TipoCambio"] = DefinedName("TipoCambio", attr_text=f"Configuracion!$C${cfg['ROW_TC']}")
wb.defined_names["MonedaVisualizacion"] = DefinedName("MonedaVisualizacion", attr_text=f"Configuracion!$C${ROW_VIS}")
wb.defined_names["FactorConversion"] = DefinedName("FactorConversion", attr_text=f"Configuracion!$C${ROW_FACTOR}")

order = ["Portada", "Instrucciones", "Configuracion", "Categorias", "Flujo de Caja", "Resumen", "Dashboard"]
wb._sheets = [wb[name] for name in order]
tab_colors = {
    "Portada": NAVY, "Instrucciones": TEAL, "Configuracion": GOLD, "Categorias": GOLD,
    "Flujo de Caja": TEAL, "Resumen": NAVY, "Dashboard": NAVY,
}
for name, color in tab_colors.items():
    wb[name].sheet_properties.tabColor = color

wb.active = wb.index(wb["Dashboard"])
wb.calculation.fullCalcOnLoad = True

for name in order:
    psheet = wb[name]
    psheet.page_setup.orientation = "landscape"
    psheet.page_setup.fitToWidth = 1
    psheet.page_setup.fitToHeight = 0
    psheet.sheet_properties.pageSetUpPr.fitToPage = True
    psheet.page_margins.left = 0.4
    psheet.page_margins.right = 0.4
    psheet.page_margins.top = 0.5
    psheet.page_margins.bottom = 0.5

out_path = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "productos", "05-flujo-de-caja", "Flujo-de-Caja-para-Negocios.xlsx",
)
os.makedirs(os.path.dirname(out_path), exist_ok=True)
wb.save(out_path)
print("Saved:", out_path)
print("Key rows:", dict(RES_TOTAL_ROW=RES_TOTAL_ROW, MONTH_FIRST=MONTH_FIRST, MONTH_LAST=MONTH_LAST))
