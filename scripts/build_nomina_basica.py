import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.workbook.defined_name import DefinedName
import datetime
import os

NAVY = "1F3B57"
TEAL = "2CA58D"
GOLD = "C9A24B"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D9DEE3"
WHITE = "FFFFFF"
DARK_TEXT = "1A1A1A"
RED_NEG = "B23B3B"
INPUT_FILL = "FFF6D9"
LOCKED_TEXT = "555555"
FONT_NAME = "Arial"

f_h1 = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
f_h2 = Font(name=FONT_NAME, size=12, bold=True, color=WHITE)
f_h3 = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
f_body = Font(name=FONT_NAME, size=10, color=DARK_TEXT)
f_body_b = Font(name=FONT_NAME, size=10, bold=True, color=DARK_TEXT)
f_kpi_label = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
f_kpi_value = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
f_note = Font(name=FONT_NAME, size=9, italic=True, color=LOCKED_TEXT)
f_input = Font(name=FONT_NAME, size=10, color="0000FF")

fill_header = PatternFill("solid", fgColor=NAVY)
fill_header2 = PatternFill("solid", fgColor=TEAL)
fill_light = PatternFill("solid", fgColor=LIGHT_GRAY)
fill_input = PatternFill("solid", fgColor=INPUT_FILL)
fill_white = PatternFill("solid", fgColor=WHITE)

thin = Side(style="thin", color=MID_GRAY)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

PROTECT_PASSWORD = "plantilla2026"

wb = Workbook()
wb.remove(wb.active)


def style_header_row(ws, row, col_start, col_end, fill=fill_header, font=f_h2, height=22):
    ws.row_dimensions[row].height = height
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align_center
        cell.border = border_all


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def lock_all(ws):
    ws.protection.sheet = True
    ws.protection.password = PROTECT_PASSWORD
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def unlock_range(ws, cell_range):
    for row in ws[cell_range]:
        for cell in row:
            cell.protection = Protection(locked=False)


def build_portada(title, subtitle, features, price):
    ws = wb.create_sheet("Portada")
    set_col_widths(ws, [4, 22, 22, 22, 22, 22, 4])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    ws["B2"] = "LÍNEA DE PRODUCTOS DIGITALES · PLANTILLAS EXCEL"
    ws["B2"].font = Font(name=FONT_NAME, size=10, bold=True, color=TEAL)
    ws.merge_cells("B4:F6")
    ws["B4"] = title
    ws["B4"].font = Font(name=FONT_NAME, size=28, bold=True, color=NAVY)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells("B7:F8")
    ws["B7"] = subtitle
    ws["B7"].font = Font(name=FONT_NAME, size=11, color=DARK_TEXT)
    ws["B7"].alignment = align_wrap
    for r in range(2, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[4].height = 30
    row0 = 10
    ws.merge_cells(f"B{row0}:F{row0}")
    ws[f"B{row0}"] = "QUÉ INCLUYE ESTA PLANTILLA"
    ws[f"B{row0}"].font = f_h3
    for c in range(2, 7):
        ws.cell(row=row0, column=c).fill = fill_light
    ws.row_dimensions[row0].height = 20
    r = row0 + 1
    for feat in features:
        ws.merge_cells(f"B{r}:F{r}")
        ws[f"B{r}"] = f"✔  {feat}"
        ws[f"B{r}"].font = f_body
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = f"Producto digital · Precio orientativo: {price}  |  Licencia de uso comercial"
    ws[f"B{r}"].font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD)
    r += 2
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = "© Tu Marca de Plantillas · www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_instrucciones(steps):
    ws = wb.create_sheet("Instrucciones")
    set_col_widths(ws, [3, 46, 46, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:C2")
    ws["B2"] = "CÓMO USAR ESTA PLANTILLA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    r = 4
    for title, body in steps:
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = title
        ws[f"B{r}"].font = f_h3
        for cc in (2, 3):
            ws.cell(row=r, column=cc).fill = fill_light
        ws.row_dimensions[r].height = 18
        r += 1
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = body
        ws[f"B{r}"].font = f_body
        ws[f"B{r}"].alignment = align_wrap
        ws.row_dimensions[r].height = 42
        r += 2
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = "Soporte y otras plantillas de esta línea de productos: www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_configuracion(extra_rows=None):
    ws = wb.create_sheet("Configuracion")
    set_col_widths(ws, [3, 32, 20, 40, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:D2")
    ws["B2"] = "CONFIGURACIÓN Y MONEDA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    headers_row = 4
    ws.cell(row=headers_row, column=2, value="Parámetro")
    ws.cell(row=headers_row, column=3, value="Valor")
    ws.cell(row=headers_row, column=4, value="Notas")
    style_header_row(ws, headers_row, 2, 4)

    rows = [
        ("Nombre de moneda local", "Soles (PEN)", "Solo referencia. Escribe el nombre de tu moneda."),
        ("Símbolo de moneda local", "S/", "Símbolo que se usará en el Dashboard cuando la vista esté en LOCAL."),
        ("Tipo de cambio (1 USD =)", 3.75, "Cuántas unidades de tu moneda local equivalen a 1 USD."),
        ("Moneda de visualización", "Local", "Elige LOCAL o USD. Afecta Resumen y Dashboard."),
    ]
    if extra_rows:
        rows.extend(extra_rows)

    row_map = {}
    r = 5
    for label, value, note in rows:
        ws.cell(row=r, column=2, value=label).font = f_body_b
        c = ws.cell(row=r, column=3, value=value)
        c.font = f_input
        c.fill = fill_input
        c.alignment = align_center
        c.border = border_all
        ws.cell(row=r, column=4, value=note).font = f_note
        ws.cell(row=r, column=4).alignment = align_wrap
        ws.row_dimensions[r].height = 18
        ws.cell(row=r, column=2).border = border_all
        row_map[label] = r
        r += 1

    ROW_VIS = row_map["Moneda de visualización"]
    ROW_TC = row_map["Tipo de cambio (1 USD =)"]
    ROW_SIMBOLO = row_map["Símbolo de moneda local"]

    r_calc_label = r
    ws.cell(row=r_calc_label, column=2, value="Valores calculados (no editar)").font = f_h3
    for cc in (2, 3, 4):
        ws.cell(row=r_calc_label, column=cc).fill = fill_light
    ws.row_dimensions[r_calc_label].height = 18
    r += 1

    ROW_FACTOR = r
    ws.cell(row=ROW_FACTOR, column=2, value="Factor de conversión aplicado").font = f_body_b
    ws.cell(row=ROW_FACTOR, column=3, value=f'=IF(C{ROW_VIS}="USD",1/C{ROW_TC},1)')
    ws.cell(row=ROW_FACTOR, column=3).number_format = "0.0000"
    ws.cell(row=ROW_FACTOR, column=3).border = border_all
    ws.cell(row=ROW_FACTOR, column=4, value="Vale 1 si vista LOCAL; vale 1/Tipo de cambio si vista USD.").font = f_note
    ws.row_dimensions[ROW_FACTOR].height = 18
    r += 1

    ROW_SIMBOLO_ACTUAL = r
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=2, value="Símbolo de moneda actual").font = f_body_b
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3, value=f'=IF(C{ROW_VIS}="USD","US$",C{ROW_SIMBOLO})')
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).border = border_all
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).alignment = align_center
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=4, value="Símbolo que verás en el Dashboard según la moneda elegida.").font = f_note
    ws.row_dimensions[ROW_SIMBOLO_ACTUAL].height = 18
    r += 1

    note_row = r + 1
    ws.merge_cells(f"B{note_row}:D{note_row}")
    ws[f"B{note_row}"] = "Cambia 'Moneda de visualización' entre Local y USD para recalcular automáticamente toda la plantilla."
    ws[f"B{note_row}"].font = f_note
    ws[f"B{note_row}"].alignment = align_wrap

    lock_all(ws)
    for label in row_map:
        rr = row_map[label]
        unlock_range(ws, f"C{rr}:C{rr}")

    dv_moneda = DataValidation(type="list", formula1='"Local,USD"', allow_blank=False)
    ws.add_data_validation(dv_moneda)
    dv_moneda.add(ws[f"C{ROW_VIS}"])

    return dict(ROW_TC=ROW_TC, ROW_VIS=ROW_VIS, ROW_SIMBOLO=ROW_SIMBOLO,
                ROW_FACTOR=ROW_FACTOR, ROW_SIMBOLO_ACTUAL=ROW_SIMBOLO_ACTUAL, row_map=row_map)


MONTHS_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
MONTHS_SHORT = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# =========================================================
build_portada(
    "NÓMINA BÁSICA",
    ("Organiza la planilla de tu equipo: sueldos, bonificaciones y descuentos, mes a mes, y "
     "obtén el sueldo neto de cada colaborador sin fórmulas complicadas — con panel de control, "
     "gráficos automáticos y selector de moneda (local / USD)."),
    [
        "Dashboard con indicadores clave (KPIs) y gráficos automáticos",
        "Fórmulas automáticas: nada que calcular a mano",
        "Cálculo automático de descuentos (seguro social e impuesto) y sueldo neto",
        "Gráficos: pastel (composición de planilla), barras y evolución mensual",
        "Resumen dinámico tipo tabla dinámica por colaborador y por mes",
        "Listas desplegables y validación de datos en cada registro",
        "Selector de moneda: local o USD, con tipo de cambio editable",
        "Hoja de instrucciones paso a paso",
        "Celdas y fórmulas protegidas: solo editas donde debes",
        "Diseño profesional y ejemplos precargados",
    ],
    "$12.99 USD",
)

build_instrucciones([
    ("1. Configuración",
     "Define el símbolo de tu moneda local, el tipo de cambio, las tasas de descuento (seguro "
     "social e impuesto/retención), el mes que quieres analizar y si la plantilla se muestra en "
     "moneda LOCAL o USD."),
    ("2. Empleados",
     "En 'Empleados' registra a cada colaborador: nombre, cargo, fecha de ingreso y sueldo base "
     "mensual. Hay espacio para hasta 10 colaboradores."),
    ("3. Bonificaciones",
     "En 'Bonificaciones' registra, mes a mes, cualquier bono, comisión u horas extra de cada "
     "colaborador (además de su sueldo base)."),
    ("4. Resumen automático",
     "La hoja 'Resumen' calcula, para el mes elegido, el total de ingresos, los descuentos y el "
     "sueldo neto de cada colaborador, y totaliza la planilla mes a mes durante el año."),
    ("5. Dashboard",
     "La hoja 'Dashboard' muestra tu planilla bruta y neta del mes, el número de colaboradores y "
     "los descuentos totales, con gráficos que se actualizan solos."),
    ("6. Cambiar de mes o de moneda",
     "En 'Configuración' cambia 'Mes seleccionado' para analizar otro mes, o 'Moneda de "
     "visualización' entre Local y USD: todo se recalcula automáticamente."),
    ("7. Celdas protegidas",
     "Las hojas están protegidas para que no borres fórmulas por accidente. Solo puedes escribir en "
     "las celdas resaltadas en color crema / texto azul. Contraseña de desprotección: plantilla2026."),
    ("Nota",
     "Las tasas de descuento son simplificadas (un % fijo sobre el total de ingresos) para que "
     "cualquier país pueda adaptarlas. No reemplazan el cálculo legal de planillas de tu país; "
     "para eso consulta a un contador o especialista en nómina."),
])

cfg = build_configuracion(extra_rows=[
    ("Tasa de seguro social (%)", 0.13, "Porcentaje de descuento por seguro social/pensión sobre el total de ingresos."),
    ("Tasa de impuesto/retención (%)", 0.08, "Porcentaje de retención de impuesto sobre el total de ingresos."),
    ("Mes seleccionado", "Enero", "Mes que se muestra en Resumen y Dashboard."),
    ("Año de análisis", 2026, "Año que se usa como referencia."),
])
ROW_FACTOR = cfg["ROW_FACTOR"]
ROW_SIMBOLO_ACTUAL = cfg["ROW_SIMBOLO_ACTUAL"]
ROW_VIS = cfg["ROW_VIS"]
ROW_TASA_SEGURO = cfg["row_map"]["Tasa de seguro social (%)"]
ROW_TASA_IMPUESTO = cfg["row_map"]["Tasa de impuesto/retención (%)"]
ROW_MES = cfg["row_map"]["Mes seleccionado"]
ROW_ANIO = cfg["row_map"]["Año de análisis"]

ws_cfg = wb["Configuracion"]
ROW_MES_NUM = ROW_ANIO + 2
ws_cfg.cell(row=ROW_MES_NUM, column=2, value="Número de mes seleccionado").font = f_body_b
ws_cfg.cell(row=ROW_MES_NUM, column=3,
            value=f'=MATCH(C{ROW_MES},{{"Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"}},0)')
ws_cfg.cell(row=ROW_MES_NUM, column=3).border = border_all
ws_cfg.cell(row=ROW_MES_NUM, column=3).alignment = align_center
ws_cfg.cell(row=ROW_MES_NUM, column=4, value="1=Enero ... 12=Diciembre.").font = f_note
ws_cfg.cell(row=ROW_MES_NUM, column=2).border = border_all
ws_cfg.row_dimensions[ROW_MES_NUM].height = 18

dv_mes_cfg = DataValidation(type="list", formula1='"' + ",".join(MONTHS_ES) + '"', allow_blank=False)
ws_cfg.add_data_validation(dv_mes_cfg)
dv_mes_cfg.add(ws_cfg[f"C{ROW_MES}"])

# =========================================================
# EMPLEADOS
# =========================================================
ws = wb.create_sheet("Empleados")
set_col_widths(ws, [10, 24, 20, 16, 16, 3])
ws.sheet_view.showGridLines = False
ws.merge_cells("A2:E2")
ws["A2"] = "PLANTILLA DE EMPLEADOS"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26

e_header = 4
e_headers = ["Código", "Nombre", "Cargo", "Fecha de ingreso", "Sueldo base mensual"]
for i, h in enumerate(e_headers, start=1):
    ws.cell(row=e_header, column=i, value=h)
style_header_row(ws, e_header, 1, 5, height=24)

E_FIRST = e_header + 1
example_employees = [
    ("EMP-01", "Ana Torres", "Gerente General", datetime.date(2023, 3, 1), 2500),
    ("EMP-02", "Luis Medina", "Contador", datetime.date(2023, 6, 15), 1600),
    ("EMP-03", "Sofía Vargas", "Ejecutiva de Ventas", datetime.date(2024, 1, 10), 1200),
    ("EMP-04", "Jorge Paredes", "Asistente Administrativo", datetime.date(2024, 4, 1), 1000),
    ("EMP-05", "Valeria Ríos", "Diseñadora", datetime.date(2025, 2, 1), 1300),
]
E_LAST = E_FIRST + 9  # 10 slots

for i in range(E_LAST - E_FIRST + 1):
    rr = E_FIRST + i
    if i < len(example_employees):
        code, name, cargo, fecha, sueldo = example_employees[i]
        ws.cell(row=rr, column=1, value=code)
        ws.cell(row=rr, column=2, value=name)
        ws.cell(row=rr, column=3, value=cargo)
        ws.cell(row=rr, column=4, value=fecha).number_format = "DD/MM/YYYY"
        ws.cell(row=rr, column=5, value=sueldo)
    for c in range(1, 6):
        cell = ws.cell(row=rr, column=c)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        if c == 5:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
        elif c == 4:
            cell.number_format = "DD/MM/YYYY"
            cell.alignment = align_center
        elif c == 1:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

lock_all(ws)
unlock_range(ws, f"A{E_FIRST}:E{E_LAST}")

EMP_NAME_RANGE = f"Empleados!$B${E_FIRST}:$B${E_LAST}"

# =========================================================
# BONIFICACIONES (grid empleado x mes)
# =========================================================
ws = wb.create_sheet("Bonificaciones")
set_col_widths(ws, [24] + [9] * 12 + [13])
ws.sheet_view.showGridLines = False
ws.freeze_panes = "B5"

ws.merge_cells("A2:N2")
ws["A2"] = "BONIFICACIONES / HORAS EXTRA (por colaborador y mes)"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("A3:N3")
ws["A3"] = "Monto adicional al sueldo base de cada colaborador en cada mes (bonos, comisiones, horas extra)."
ws["A3"].font = f_note

b_header = 4
ws.cell(row=b_header, column=1, value="Colaborador")
for i, m in enumerate(MONTHS_SHORT):
    ws.cell(row=b_header, column=2 + i, value=m)
ws.cell(row=b_header, column=14, value="Total anual")
style_header_row(ws, b_header, 1, 14, height=20)

B_FIRST = b_header + 1
B_LAST = B_FIRST + (E_LAST - E_FIRST)

example_bonus = {
    "Ana Torres": [200, 0, 300, 0, 0, 200, 0, 0, 300, 0, 0, 500],
    "Luis Medina": [0, 0, 100, 0, 0, 0, 0, 0, 100, 0, 0, 200],
    "Sofía Vargas": [150, 220, 180, 0, 0, 0, 0, 0, 0, 0, 0, 300],
    "Jorge Paredes": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 150],
    "Valeria Ríos": [0, 100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 200],
}

for i in range(E_LAST - E_FIRST + 1):
    rr = B_FIRST + i
    e_rr = E_FIRST + i
    ws.cell(row=rr, column=1, value=f"=Empleados!$B${e_rr}")
    ws.cell(row=rr, column=1).font = f_body
    ws.cell(row=rr, column=1).border = border_all

    emp_name = example_employees[i][1] if i < len(example_employees) else None
    vals = example_bonus.get(emp_name, [None] * 12)

    for m in range(12):
        cell = ws.cell(row=rr, column=2 + m)
        val = vals[m] if m < len(vals) else None
        if val is not None:
            cell.value = val
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        cell.number_format = "#,##0"
        cell.alignment = align_center

    tot = ws.cell(row=rr, column=14, value=f"=SUM(B{rr}:M{rr})")
    tot.number_format = "#,##0"
    tot.font = f_body_b
    tot.fill = fill_light
    tot.border = border_all
    tot.alignment = align_center

dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                         allow_blank=True, showErrorMessage=True,
                         errorTitle="Monto inválido", error="Ingresa un número igual o mayor a 0.")
ws.add_data_validation(dv_num)
dv_num.add(f"B{B_FIRST}:M{B_LAST}")

lock_all(ws)
unlock_range(ws, f"B{B_FIRST}:M{B_LAST}")

# =========================================================
# RESUMEN
# =========================================================
ws = wb.create_sheet("Resumen")
set_col_widths(ws, [3, 22, 20, 14, 14, 14, 14, 14, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:H2")
ws["B2"] = "RESUMEN DINÁMICO (se actualiza solo)"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("B3:H3")
ws["B3"] = f'=" Mes: " & Configuracion!$C${ROW_MES} & "   |   Moneda: " & Configuracion!$C${ROW_VIS}'
ws["B3"].font = f_note

r0 = 5
ws.cell(row=r0, column=2, value="PLANILLA POR COLABORADOR (mes seleccionado)")
ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=8)
ws.cell(row=r0, column=2).font = f_h2
for c in range(2, 9):
    ws.cell(row=r0, column=c).fill = fill_header2
ws.row_dimensions[r0].height = 20

hr = r0 + 1
res_headers = ["Colaborador", "Cargo", "Sueldo base", "Bonificación", "Total ingresos", "Descuentos", "Sueldo neto"]
for i, h in enumerate(res_headers):
    ws.cell(row=hr, column=2 + i, value=h)
style_header_row(ws, hr, 2, 8, height=26)

RES_FIRST = hr + 1
RES_LAST = RES_FIRST + (E_LAST - E_FIRST)

for i in range(E_LAST - E_FIRST + 1):
    rr = RES_FIRST + i
    e_rr = E_FIRST + i
    b_rr = B_FIRST + i
    ws.cell(row=rr, column=2, value=f'=IF(Empleados!$A${e_rr}="","",Empleados!$B${e_rr})')
    ws.cell(row=rr, column=3, value=f'=IF(Empleados!$A${e_rr}="","",Empleados!$C${e_rr})')
    ws.cell(row=rr, column=4,
            value=f'=IF(Empleados!$A${e_rr}="",0,Empleados!$E${e_rr}*Configuracion!$C${ROW_FACTOR})')
    ws.cell(row=rr, column=5,
            value=(f'=IF(Empleados!$A${e_rr}="",0,'
                   f'INDEX(Bonificaciones!$B${b_rr}:$M${b_rr},Configuracion!$C${ROW_MES_NUM})'
                   f'*Configuracion!$C${ROW_FACTOR})'))
    ws.cell(row=rr, column=6, value=f"=D{rr}+E{rr}")
    ws.cell(row=rr, column=7,
            value=f'=F{rr}*(Configuracion!$C${ROW_TASA_SEGURO}+Configuracion!$C${ROW_TASA_IMPUESTO})')
    ws.cell(row=rr, column=8, value=f"=F{rr}-G{rr}")
    for c in range(2, 9):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c in (4, 5, 6, 7, 8):
            cell.number_format = "#,##0.00"
            cell.alignment = align_center

tot_row = RES_LAST + 1
ws.cell(row=tot_row, column=2, value="TOTAL").font = f_body_b
for c, col_letter in zip(range(4, 9), ["D", "E", "F", "G", "H"]):
    cell = ws.cell(row=tot_row, column=c, value=f"=SUM({col_letter}{RES_FIRST}:{col_letter}{RES_LAST})")
    cell.number_format = "#,##0.00"
    cell.font = f_body_b
    cell.fill = fill_light
    cell.alignment = align_center
for c in range(2, 9):
    ws.cell(row=tot_row, column=c).border = border_all
    ws.cell(row=tot_row, column=c).fill = fill_light
RES_TOTAL_ROW = tot_row

n_emp_row = tot_row + 1
ws.cell(row=n_emp_row, column=2, value="N° de colaboradores activos").font = f_body_b
ws.cell(row=n_emp_row, column=2).border = border_all
ws.cell(row=n_emp_row, column=4, value=f'=COUNTIF(Empleados!$A${E_FIRST}:$A${E_LAST},"<>")')
ws.cell(row=n_emp_row, column=4).border = border_all
ws.cell(row=n_emp_row, column=4).alignment = align_center
ws.cell(row=n_emp_row, column=4).font = f_body_b
N_EMP_ROW = n_emp_row

# ---- Evolución mensual de la planilla ----
r1 = n_emp_row + 3
ws.cell(row=r1, column=2, value="EVOLUCIÓN MENSUAL DE LA PLANILLA")
ws.merge_cells(start_row=r1, start_column=2, end_row=r1, end_column=5)
ws.cell(row=r1, column=2).font = f_h2
for c in range(2, 6):
    ws.cell(row=r1, column=c).fill = fill_header2
ws.row_dimensions[r1].height = 20

hr2 = r1 + 1
mon_headers = ["Mes", "Planilla bruta", "Descuentos", "Planilla neta"]
for i, h in enumerate(mon_headers):
    ws.cell(row=hr2, column=2 + i, value=h)
style_header_row(ws, hr2, 2, 5, height=18)

MONTH_FIRST = hr2 + 1
for i, mname in enumerate(MONTHS_ES):
    rr = MONTH_FIRST + i
    col_letter = get_column_letter(2 + i)
    ws.cell(row=rr, column=2, value=mname)
    ws.cell(row=rr, column=3,
            value=(f"=(SUM(Empleados!$E${E_FIRST}:$E${E_LAST})+SUM(Bonificaciones!{col_letter}${B_FIRST}:{col_letter}${B_LAST}))"
                   f"*Configuracion!$C${ROW_FACTOR}"))
    ws.cell(row=rr, column=4,
            value=f'=C{rr}*(Configuracion!$C${ROW_TASA_SEGURO}+Configuracion!$C${ROW_TASA_IMPUESTO})')
    ws.cell(row=rr, column=5, value=f"=C{rr}-D{rr}")
    for c in (2, 3, 4, 5):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c in (3, 4, 5):
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
MONTH_LAST = MONTH_FIRST + len(MONTHS_ES) - 1

lock_all(ws)

# =========================================================
# DASHBOARD
# =========================================================
ws = wb.create_sheet("Dashboard")
set_col_widths(ws, [3, 16, 16, 16, 16, 16, 16, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:H2")
ws["B2"] = "DASHBOARD · NÓMINA BÁSICA"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 28
ws.merge_cells("B3:H3")
ws["B3"] = f'=" Mes: " & Configuracion!$C${ROW_MES} & "   |   Moneda: " & Configuracion!$C${ROW_VIS}'
ws["B3"].font = f_note

kpi_row = 5
card_defs = [
    ("PLANILLA BRUTA (MES)", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!F{RES_TOTAL_ROW},\"#,##0.00\")", TEAL),
    ("DESCUENTOS (MES)", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!G{RES_TOTAL_ROW},\"#,##0.00\")", RED_NEG),
    ("PLANILLA NETA (MES)", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!H{RES_TOTAL_ROW},\"#,##0.00\")", NAVY),
    ("N° COLABORADORES", f"=Resumen!D{N_EMP_ROW}", GOLD),
]
col_starts = [2, 4, 6, 8]
for (label, formula, color), col0 in zip(card_defs, col_starts):
    ws.merge_cells(start_row=kpi_row, start_column=col0, end_row=kpi_row, end_column=col0 + 1)
    lc = ws.cell(row=kpi_row, column=col0, value=label)
    lc.font = f_kpi_label
    lc.alignment = align_center
    for c in range(col0, col0 + 2):
        ws.cell(row=kpi_row, column=c).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=kpi_row + 1, start_column=col0, end_row=kpi_row + 2, end_column=col0 + 1)
    vc = ws.cell(row=kpi_row + 1, column=col0, value=formula)
    vc.font = f_kpi_value
    vc.alignment = align_center
    for rr in (kpi_row + 1, kpi_row + 2):
        for c in range(col0, col0 + 2):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor=color)

ws.row_dimensions[kpi_row].height = 18
ws.row_dimensions[kpi_row + 1].height = 22
ws.row_dimensions[kpi_row + 2].height = 22

chart_top = kpi_row + 4

pie = PieChart()
pie.title = "Sueldo neto por colaborador (mes)"
data_pie = Reference(wb["Resumen"], min_col=8, min_row=hr, max_row=RES_LAST)
cats_pie = Reference(wb["Resumen"], min_col=2, min_row=RES_FIRST, max_row=RES_LAST)
pie.add_data(data_pie, titles_from_data=True)
pie.set_categories(cats_pie)
pie.height = 9
pie.width = 15
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showVal = False
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showLegendKey = False
pie.dataLabels.showBubbleSize = False
ws.add_chart(pie, f"B{chart_top}")

bar = BarChart()
bar.type = "col"
bar.title = "Ingresos vs. sueldo neto por colaborador"
bar.style = 10
bdata = Reference(wb["Resumen"], min_col=6, max_col=8, min_row=hr, max_row=RES_LAST)
bcats = Reference(wb["Resumen"], min_col=2, min_row=RES_FIRST, max_row=RES_LAST)
bar.add_data(bdata, titles_from_data=True)
bar.set_categories(bcats)
bar.height = 9
bar.width = 15
ws.add_chart(bar, f"E{chart_top}")

line = LineChart()
line.title = "Evolución mensual de la planilla neta"
ldata = Reference(wb["Resumen"], min_col=5, min_row=hr2, max_row=MONTH_LAST)
lcats = Reference(wb["Resumen"], min_col=2, min_row=MONTH_FIRST, max_row=MONTH_LAST)
line.add_data(ldata, titles_from_data=True)
line.set_categories(lcats)
line.height = 9
line.width = 30
ws.add_chart(line, f"B{chart_top + 19}")

lock_all(ws)

# =========================================================
wb.defined_names["TipoCambio"] = DefinedName("TipoCambio", attr_text=f"Configuracion!$C${cfg['ROW_TC']}")
wb.defined_names["MonedaVisualizacion"] = DefinedName("MonedaVisualizacion", attr_text=f"Configuracion!$C${ROW_VIS}")
wb.defined_names["FactorConversion"] = DefinedName("FactorConversion", attr_text=f"Configuracion!$C${ROW_FACTOR}")

order = ["Portada", "Instrucciones", "Configuracion", "Empleados", "Bonificaciones", "Resumen", "Dashboard"]
wb._sheets = [wb[name] for name in order]
tab_colors = {
    "Portada": NAVY, "Instrucciones": TEAL, "Configuracion": GOLD, "Empleados": GOLD,
    "Bonificaciones": TEAL, "Resumen": NAVY, "Dashboard": NAVY,
}
for name, color in tab_colors.items():
    wb[name].sheet_properties.tabColor = color

wb.active = wb.index(wb["Dashboard"])
wb.calculation.fullCalcOnLoad = True

for name in order:
    psheet = wb[name]
    psheet.page_setup.orientation = "landscape"
    psheet.page_setup.fitToWidth = 1
    psheet.page_setup.fitToHeight = 0
    psheet.sheet_properties.pageSetUpPr.fitToPage = True
    psheet.page_margins.left = 0.4
    psheet.page_margins.right = 0.4
    psheet.page_margins.top = 0.5
    psheet.page_margins.bottom = 0.5

out_path = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "productos", "08-nomina-basica", "Nomina-Basica.xlsx",
)
os.makedirs(os.path.dirname(out_path), exist_ok=True)
wb.save(out_path)
print("Saved:", out_path)
print("Key rows:", dict(RES_TOTAL_ROW=RES_TOTAL_ROW, N_EMP_ROW=N_EMP_ROW, ROW_MES_NUM=ROW_MES_NUM))
