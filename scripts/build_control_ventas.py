import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.defined_name import DefinedName
import datetime
import os

NAVY = "1F3B57"
TEAL = "2CA58D"
GOLD = "C9A24B"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D9DEE3"
WHITE = "FFFFFF"
DARK_TEXT = "1A1A1A"
RED_NEG = "B23B3B"
INPUT_FILL = "FFF6D9"
LOCKED_TEXT = "555555"
FONT_NAME = "Arial"

f_h1 = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
f_h2 = Font(name=FONT_NAME, size=12, bold=True, color=WHITE)
f_h3 = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
f_body = Font(name=FONT_NAME, size=10, color=DARK_TEXT)
f_body_b = Font(name=FONT_NAME, size=10, bold=True, color=DARK_TEXT)
f_kpi_label = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
f_kpi_value = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
f_note = Font(name=FONT_NAME, size=9, italic=True, color=LOCKED_TEXT)
f_input = Font(name=FONT_NAME, size=10, color="0000FF")

fill_header = PatternFill("solid", fgColor=NAVY)
fill_header2 = PatternFill("solid", fgColor=TEAL)
fill_light = PatternFill("solid", fgColor=LIGHT_GRAY)
fill_input = PatternFill("solid", fgColor=INPUT_FILL)
fill_white = PatternFill("solid", fgColor=WHITE)

thin = Side(style="thin", color=MID_GRAY)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)


wb = Workbook()
wb.remove(wb.active)


def style_header_row(ws, row, col_start, col_end, fill=fill_header, font=f_h2, height=22):
    ws.row_dimensions[row].height = height
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align_center
        cell.border = border_all


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def lock_all(ws):
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def unlock_range(ws, cell_range):
    for row in ws[cell_range]:
        for cell in row:
            cell.protection = Protection(locked=False)


def build_portada(title, subtitle, features, price):
    ws = wb.create_sheet("Portada")
    set_col_widths(ws, [4, 22, 22, 22, 22, 22, 4])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    ws["B2"] = "LÍNEA DE PRODUCTOS DIGITALES · PLANTILLAS EXCEL"
    ws["B2"].font = Font(name=FONT_NAME, size=10, bold=True, color=TEAL)
    ws.merge_cells("B4:F6")
    ws["B4"] = title
    ws["B4"].font = Font(name=FONT_NAME, size=28, bold=True, color=NAVY)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells("B7:F8")
    ws["B7"] = subtitle
    ws["B7"].font = Font(name=FONT_NAME, size=11, color=DARK_TEXT)
    ws["B7"].alignment = align_wrap
    for r in range(2, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[4].height = 30
    row0 = 10
    ws.merge_cells(f"B{row0}:F{row0}")
    ws[f"B{row0}"] = "QUÉ INCLUYE ESTA PLANTILLA"
    ws[f"B{row0}"].font = f_h3
    for c in range(2, 7):
        ws.cell(row=row0, column=c).fill = fill_light
    ws.row_dimensions[row0].height = 20
    r = row0 + 1
    for feat in features:
        ws.merge_cells(f"B{r}:F{r}")
        ws[f"B{r}"] = f"✔  {feat}"
        ws[f"B{r}"].font = f_body
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = f"Producto digital · Precio orientativo: {price}  |  Licencia de uso comercial"
    ws[f"B{r}"].font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD)
    r += 2
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = "© Tu Marca de Plantillas · www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_instrucciones(steps):
    ws = wb.create_sheet("Instrucciones")
    set_col_widths(ws, [3, 46, 46, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:C2")
    ws["B2"] = "CÓMO USAR ESTA PLANTILLA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    r = 4
    for title, body in steps:
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = title
        ws[f"B{r}"].font = f_h3
        for cc in (2, 3):
            ws.cell(row=r, column=cc).fill = fill_light
        ws.row_dimensions[r].height = 18
        r += 1
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = body
        ws[f"B{r}"].font = f_body
        ws[f"B{r}"].alignment = align_wrap
        ws.row_dimensions[r].height = 42
        r += 2
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = "Soporte y otras plantillas de esta línea de productos: www.tumarca.com"
    ws[f"B{r}"].font = f_note
    lock_all(ws)


def build_configuracion(extra_rows=None):
    ws = wb.create_sheet("Configuracion")
    set_col_widths(ws, [3, 32, 20, 40, 3])
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:D2")
    ws["B2"] = "CONFIGURACIÓN Y MONEDA"
    ws["B2"].font = f_h1
    ws.row_dimensions[2].height = 26
    headers_row = 4
    ws.cell(row=headers_row, column=2, value="Parámetro")
    ws.cell(row=headers_row, column=3, value="Valor")
    ws.cell(row=headers_row, column=4, value="Notas")
    style_header_row(ws, headers_row, 2, 4)

    rows = [
        ("Nombre de moneda local", "Soles (PEN)", "Solo referencia. Escribe el nombre de tu moneda."),
        ("Símbolo de moneda local", "S/", "Símbolo que se usará en el Dashboard cuando la vista esté en LOCAL."),
        ("Tipo de cambio (1 USD =)", 3.75, "Cuántas unidades de tu moneda local equivalen a 1 USD."),
        ("Moneda de visualización", "Local", "Elige LOCAL o USD. Afecta Resumen y Dashboard."),
    ]
    if extra_rows:
        rows.extend(extra_rows)

    row_map = {}
    r = 5
    for label, value, note in rows:
        ws.cell(row=r, column=2, value=label).font = f_body_b
        c = ws.cell(row=r, column=3, value=value)
        c.font = f_input
        c.fill = fill_input
        c.alignment = align_center
        c.border = border_all
        ws.cell(row=r, column=4, value=note).font = f_note
        ws.cell(row=r, column=4).alignment = align_wrap
        ws.row_dimensions[r].height = 18
        ws.cell(row=r, column=2).border = border_all
        row_map[label] = r
        r += 1

    ROW_VIS = row_map["Moneda de visualización"]
    ROW_TC = row_map["Tipo de cambio (1 USD =)"]
    ROW_SIMBOLO = row_map["Símbolo de moneda local"]

    r_calc_label = r
    ws.cell(row=r_calc_label, column=2, value="Valores calculados (no editar)").font = f_h3
    for cc in (2, 3, 4):
        ws.cell(row=r_calc_label, column=cc).fill = fill_light
    ws.row_dimensions[r_calc_label].height = 18
    r += 1

    ROW_FACTOR = r
    ws.cell(row=ROW_FACTOR, column=2, value="Factor de conversión aplicado").font = f_body_b
    ws.cell(row=ROW_FACTOR, column=3, value=f'=IF(C{ROW_VIS}="USD",1/C{ROW_TC},1)')
    ws.cell(row=ROW_FACTOR, column=3).number_format = "0.0000"
    ws.cell(row=ROW_FACTOR, column=3).border = border_all
    ws.cell(row=ROW_FACTOR, column=4, value="Vale 1 si vista LOCAL; vale 1/Tipo de cambio si vista USD.").font = f_note
    ws.row_dimensions[ROW_FACTOR].height = 18
    r += 1

    ROW_SIMBOLO_ACTUAL = r
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=2, value="Símbolo de moneda actual").font = f_body_b
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3, value=f'=IF(C{ROW_VIS}="USD","US$",C{ROW_SIMBOLO})')
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).border = border_all
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=3).alignment = align_center
    ws.cell(row=ROW_SIMBOLO_ACTUAL, column=4, value="Símbolo que verás en el Dashboard según la moneda elegida.").font = f_note
    ws.row_dimensions[ROW_SIMBOLO_ACTUAL].height = 18
    r += 1

    note_row = r + 1
    ws.merge_cells(f"B{note_row}:D{note_row}")
    ws[f"B{note_row}"] = "Cambia 'Moneda de visualización' entre Local y USD para recalcular automáticamente toda la plantilla."
    ws[f"B{note_row}"].font = f_note
    ws[f"B{note_row}"].alignment = align_wrap

    lock_all(ws)
    for label in row_map:
        rr = row_map[label]
        unlock_range(ws, f"C{rr}:C{rr}")

    dv_moneda = DataValidation(type="list", formula1='"Local,USD"', allow_blank=False)
    ws.add_data_validation(dv_moneda)
    dv_moneda.add(ws[f"C{ROW_VIS}"])

    return dict(ROW_TC=ROW_TC, ROW_VIS=ROW_VIS, ROW_SIMBOLO=ROW_SIMBOLO,
                ROW_FACTOR=ROW_FACTOR, ROW_SIMBOLO_ACTUAL=ROW_SIMBOLO_ACTUAL, row_map=row_map)


MONTHS_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

# =========================================================
build_portada(
    "CONTROL DE VENTAS",
    ("Registra cada venta de tu negocio, controla qué productos se venden más, quiénes son tus "
     "mejores clientes y cuánto tienes pendiente de cobro — con panel de control, gráficos "
     "automáticos y selector de moneda (local / USD)."),
    [
        "Dashboard con indicadores clave (KPIs) y gráficos automáticos",
        "Fórmulas automáticas: nada que calcular a mano",
        "Gráficos: pastel (ventas por producto), barras (ventas por mes y top clientes)",
        "Resumen dinámico tipo tabla dinámica por producto, cliente y mes",
        "Control de cobros: pagado vs. pendiente",
        "Listas desplegables y validación de datos en cada registro",
        "Selector de moneda: local o USD, con tipo de cambio editable",
        "Hoja de instrucciones paso a paso",
        "Celdas y fórmulas protegidas: solo editas donde debes",
        "Diseño profesional y ejemplos precargados",
    ],
    "$9.99 USD",
)

build_instrucciones([
    ("1. Configuración",
     "Define el símbolo de tu moneda local, el tipo de cambio y si la plantilla se muestra en "
     "moneda LOCAL o USD."),
    ("2. Catálogo y Clientes",
     "En 'Catálogo' registra tus productos/servicios y su precio. En 'Clientes' registra tu "
     "cartera de clientes. Ambos alimentan las listas desplegables de 'Ventas'."),
    ("3. Ventas",
     "En 'Ventas' registra cada venta: fecha, cliente (lista desplegable), producto (lista "
     "desplegable), cantidad, método de pago y estado (Pagado/Pendiente). El precio unitario y "
     "el total se calculan solos."),
    ("4. Resumen automático",
     "La hoja 'Resumen' agrupa automáticamente tus ventas por producto, por cliente y por mes, "
     "como una tabla dinámica."),
    ("5. Dashboard",
     "La hoja 'Dashboard' muestra tus ventas del año, tu ticket promedio, lo pendiente de cobro y "
     "tu producto más vendido, con gráficos que se actualizan solos."),
    ("6. Cambiar de moneda",
     "Cambia 'Moneda de visualización' en 'Configuración' entre Local y USD: todo se recalcula "
     "automáticamente."),
    ("7. Celdas protegidas",
     "Las hojas están protegidas para que no borres fórmulas por accidente. Solo puedes escribir en "
     "las celdas resaltadas en color crema / texto azul. Si necesitas editar otra celda, ve a Revisar → Desproteger hoja (no pide contraseña)."),
])

cfg = build_configuracion(extra_rows=[
    ("Año de análisis", 2026, "Año que se usa para agrupar el Resumen Mensual."),
])
ROW_FACTOR = cfg["ROW_FACTOR"]
ROW_SIMBOLO_ACTUAL = cfg["ROW_SIMBOLO_ACTUAL"]
ROW_VIS = cfg["ROW_VIS"]
ROW_ANIO = cfg["row_map"]["Año de análisis"]

# =========================================================
# CATALOGO
# =========================================================
ws = wb.create_sheet("Catalogo")
set_col_widths(ws, [26, 20, 16, 3])
ws.sheet_view.showGridLines = False
ws.merge_cells("A2:C2")
ws["A2"] = "CATÁLOGO DE PRODUCTOS / SERVICIOS"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26

cat_header = 4
ws.cell(row=cat_header, column=1, value="Producto / Servicio")
ws.cell(row=cat_header, column=2, value="Categoría")
ws.cell(row=cat_header, column=3, value="Precio unitario")
style_header_row(ws, cat_header, 1, 3, height=20)

CAT_FIRST = cat_header + 1
example_products = [
    ("Consultoría por hora", "Servicios", 45),
    ("Diseño de logo", "Servicios", 150),
    ("Plantilla Excel básica", "Productos digitales", 7),
    ("Plantilla Excel avanzada", "Productos digitales", 15),
    ("Paquete de mantenimiento mensual", "Servicios", 200),
    ("Curso en línea", "Productos digitales", 40),
]
CAT_LAST = CAT_FIRST + 7  # 8 slots (keeps pie legend readable)

categorias_prod = ["Servicios", "Productos digitales", "Productos físicos", "Otros"]

for i in range(CAT_LAST - CAT_FIRST + 1):
    rr = CAT_FIRST + i
    if i < len(example_products):
        name, cat, price = example_products[i]
        ws.cell(row=rr, column=1, value=name)
        ws.cell(row=rr, column=2, value=cat)
        ws.cell(row=rr, column=3, value=price)
    for c in range(1, 4):
        cell = ws.cell(row=rr, column=c)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        if c == 3:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
        else:
            cell.alignment = align_left

dv_catprod = DataValidation(type="list", formula1='"' + ",".join(categorias_prod) + '"', allow_blank=True)
ws.add_data_validation(dv_catprod)
dv_catprod.add(f"B{CAT_FIRST}:B{CAT_LAST}")

lock_all(ws)
unlock_range(ws, f"A{CAT_FIRST}:C{CAT_LAST}")
CATALOGO_NAME_RANGE = f"Catalogo!$A${CAT_FIRST}:$A${CAT_LAST}"
CATALOGO_PRICE_RANGE = f"Catalogo!$C${CAT_FIRST}:$C${CAT_LAST}"

# =========================================================
# CLIENTES
# =========================================================
ws = wb.create_sheet("Clientes")
set_col_widths(ws, [26, 20, 3])
ws.sheet_view.showGridLines = False
ws.merge_cells("A2:B2")
ws["A2"] = "CARTERA DE CLIENTES"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26

cli_header = 4
ws.cell(row=cli_header, column=1, value="Cliente")
ws.cell(row=cli_header, column=2, value="Contacto")
style_header_row(ws, cli_header, 1, 2, height=20)

CLI_FIRST = cli_header + 1
example_clients = [
    ("María Fernández", "maria.f@email.com"),
    ("Carlos Rojas", "carlos.rojas@email.com"),
    ("Tienda El Sol S.A.C.", "ventas@elsol.com"),
    ("Laura Gómez", "laura.gomez@email.com"),
    ("Distribuidora Andina", "contacto@andina.com"),
]
CLI_LAST = CLI_FIRST + 6  # 7 slots (keeps bar chart readable)

for i in range(CLI_LAST - CLI_FIRST + 1):
    rr = CLI_FIRST + i
    if i < len(example_clients):
        name, contact = example_clients[i]
        ws.cell(row=rr, column=1, value=name)
        ws.cell(row=rr, column=2, value=contact)
    for c in (1, 2):
        cell = ws.cell(row=rr, column=c)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        cell.alignment = align_left

lock_all(ws)
unlock_range(ws, f"A{CLI_FIRST}:B{CLI_LAST}")
CLIENTE_RANGE = f"Clientes!$A${CLI_FIRST}:$A${CLI_LAST}"

# =========================================================
# VENTAS
# =========================================================
ws = wb.create_sheet("Ventas")
set_col_widths(ws, [12, 24, 24, 10, 14, 14, 16, 12])
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"

ws.merge_cells("A2:H2")
ws["A2"] = "REGISTRO DE VENTAS"
ws["A2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("A3:H3")
ws["A3"] = "Completa las columnas A, B, C, D, G y H. El precio unitario y el total se calculan automáticamente."
ws["A3"].font = f_note

V_HEADER = 4
v_headers = ["Fecha", "Cliente", "Producto / Servicio", "Cantidad", "Precio unitario", "Total",
             "Método de pago", "Estado"]
for i, h in enumerate(v_headers, start=1):
    ws.cell(row=V_HEADER, column=i, value=h)
style_header_row(ws, V_HEADER, 1, 8, height=26)

V_FIRST = V_HEADER + 1
V_LAST = V_FIRST + 199

metodos_pago = ["Efectivo", "Tarjeta", "Transferencia"]
estados_venta = ["Pagado", "Pendiente"]

example_sales = [
    (datetime.date(2026, 1, 5), "María Fernández", "Plantilla Excel básica", 1, "Transferencia", "Pagado"),
    (datetime.date(2026, 1, 8), "Carlos Rojas", "Consultoría por hora", 3, "Tarjeta", "Pagado"),
    (datetime.date(2026, 1, 12), "Tienda El Sol S.A.C.", "Diseño de logo", 1, "Transferencia", "Pendiente"),
    (datetime.date(2026, 1, 18), "Laura Gómez", "Plantilla Excel avanzada", 2, "Efectivo", "Pagado"),
    (datetime.date(2026, 1, 25), "Distribuidora Andina", "Paquete de mantenimiento mensual", 1, "Transferencia", "Pagado"),
    (datetime.date(2026, 2, 2), "María Fernández", "Curso en línea", 1, "Tarjeta", "Pagado"),
    (datetime.date(2026, 2, 6), "Carlos Rojas", "Plantilla Excel básica", 3, "Efectivo", "Pagado"),
    (datetime.date(2026, 2, 14), "Tienda El Sol S.A.C.", "Consultoría por hora", 5, "Transferencia", "Pendiente"),
    (datetime.date(2026, 2, 20), "Laura Gómez", "Diseño de logo", 1, "Tarjeta", "Pagado"),
    (datetime.date(2026, 3, 3), "Distribuidora Andina", "Paquete de mantenimiento mensual", 1, "Transferencia", "Pagado"),
    (datetime.date(2026, 3, 9), "María Fernández", "Plantilla Excel avanzada", 1, "Efectivo", "Pagado"),
    (datetime.date(2026, 3, 15), "Carlos Rojas", "Curso en línea", 2, "Tarjeta", "Pendiente"),
]

for i in range(V_LAST - V_FIRST + 1):
    rr = V_FIRST + i
    if i < len(example_sales):
        fecha, cliente, prod, cant, met, est = example_sales[i]
        ws.cell(row=rr, column=1, value=fecha).number_format = "DD/MM/YYYY"
        ws.cell(row=rr, column=2, value=cliente)
        ws.cell(row=rr, column=3, value=prod)
        ws.cell(row=rr, column=4, value=cant)
        ws.cell(row=rr, column=7, value=met)
        ws.cell(row=rr, column=8, value=est)
    else:
        ws.cell(row=rr, column=1).number_format = "DD/MM/YYYY"

    for col in (1, 2, 3, 4, 7, 8):
        cell = ws.cell(row=rr, column=col)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = border_all
        if col == 4:
            cell.number_format = "#,##0"
            cell.alignment = align_center
        elif col == 1:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

    pcell = ws.cell(row=rr, column=5,
                     value=f'=IF(C{rr}="","",IFERROR(INDEX({CATALOGO_PRICE_RANGE},MATCH(C{rr},{CATALOGO_NAME_RANGE},0)),0))')
    pcell.font = f_body
    pcell.fill = fill_light
    pcell.border = border_all
    pcell.number_format = "#,##0.00"
    pcell.alignment = align_center

    tcell = ws.cell(row=rr, column=6,
                     value=f'=IF(OR(C{rr}="",D{rr}=""),"",D{rr}*E{rr})')
    tcell.font = f_body
    tcell.fill = fill_light
    tcell.border = border_all
    tcell.number_format = "#,##0.00"
    tcell.alignment = align_center

dv_cliente = DataValidation(type="list", formula1=f"={CLIENTE_RANGE}", allow_blank=True)
ws.add_data_validation(dv_cliente)
dv_cliente.add(f"B{V_FIRST}:B{V_LAST}")

dv_producto = DataValidation(type="list", formula1=f"={CATALOGO_NAME_RANGE}", allow_blank=True)
ws.add_data_validation(dv_producto)
dv_producto.add(f"C{V_FIRST}:C{V_LAST}")

dv_cant = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Cantidad inválida", error="Ingresa un número igual o mayor a 0.")
ws.add_data_validation(dv_cant)
dv_cant.add(f"D{V_FIRST}:D{V_LAST}")

dv_metodo = DataValidation(type="list", formula1='"' + ",".join(metodos_pago) + '"', allow_blank=True)
ws.add_data_validation(dv_metodo)
dv_metodo.add(f"G{V_FIRST}:G{V_LAST}")

dv_estado = DataValidation(type="list", formula1='"' + ",".join(estados_venta) + '"', allow_blank=True)
ws.add_data_validation(dv_estado)
dv_estado.add(f"H{V_FIRST}:H{V_LAST}")

lock_all(ws)
unlock_range(ws, f"A{V_FIRST}:D{V_LAST}")
unlock_range(ws, f"G{V_FIRST}:H{V_LAST}")

V_FECHA_RANGE = f"Ventas!$A${V_FIRST}:$A${V_LAST}"
V_CLIENTE_RANGE = f"Ventas!$B${V_FIRST}:$B${V_LAST}"
V_PRODUCTO_RANGE = f"Ventas!$C${V_FIRST}:$C${V_LAST}"
V_CANTIDAD_RANGE = f"Ventas!$D${V_FIRST}:$D${V_LAST}"
V_TOTAL_RANGE = f"Ventas!$F${V_FIRST}:$F${V_LAST}"
V_ESTADO_RANGE = f"Ventas!$H${V_FIRST}:$H${V_LAST}"

# =========================================================
# RESUMEN
# =========================================================
ws = wb.create_sheet("Resumen")
set_col_widths(ws, [3, 26, 16, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:D2")
ws["B2"] = "RESUMEN DINÁMICO (se actualiza solo)"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 26
ws.merge_cells("B3:D3")
ws["B3"] = '=" Moneda: " & Configuracion!$C$8'
ws["B3"].font = f_note

r0 = 5
ws.cell(row=r0, column=2, value="VENTAS POR PRODUCTO")
ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=4)
ws.cell(row=r0, column=2).font = f_h2
for c in range(2, 5):
    ws.cell(row=r0, column=c).fill = fill_header2
ws.row_dimensions[r0].height = 20

hr = r0 + 1
ws.cell(row=hr, column=2, value="Producto / Servicio")
ws.cell(row=hr, column=3, value="Unidades vendidas")
ws.cell(row=hr, column=4, value="Total ventas")
style_header_row(ws, hr, 2, 4, height=26)

PROD_SUMMARY_FIRST = hr + 1
for i in range(CAT_LAST - CAT_FIRST + 1):
    rr = PROD_SUMMARY_FIRST + i
    cat_rr = CAT_FIRST + i
    ws.cell(row=rr, column=2, value=f"=Catalogo!$A${cat_rr}")
    ws.cell(row=rr, column=3, value=f'=SUMIFS({V_CANTIDAD_RANGE},{V_PRODUCTO_RANGE},$B{rr})')
    ws.cell(row=rr, column=4, value=f'=SUMIFS({V_TOTAL_RANGE},{V_PRODUCTO_RANGE},$B{rr})*Configuracion!$C${ROW_FACTOR}')
    for c in range(2, 5):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c in (3, 4):
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
PROD_SUMMARY_LAST = PROD_SUMMARY_FIRST + (CAT_LAST - CAT_FIRST)

# ---- Ventas por cliente ----
r1 = PROD_SUMMARY_LAST + 3
ws.cell(row=r1, column=2, value="VENTAS POR CLIENTE")
ws.merge_cells(start_row=r1, start_column=2, end_row=r1, end_column=3)
ws.cell(row=r1, column=2).font = f_h2
for c in range(2, 4):
    ws.cell(row=r1, column=c).fill = fill_header2
ws.row_dimensions[r1].height = 20

hr2 = r1 + 1
ws.cell(row=hr2, column=2, value="Cliente")
ws.cell(row=hr2, column=3, value="Total comprado")
style_header_row(ws, hr2, 2, 3, height=18)

CLI_SUMMARY_FIRST = hr2 + 1
for i in range(CLI_LAST - CLI_FIRST + 1):
    rr = CLI_SUMMARY_FIRST + i
    cli_rr = CLI_FIRST + i
    ws.cell(row=rr, column=2, value=f"=Clientes!$A${cli_rr}")
    ws.cell(row=rr, column=3, value=f'=SUMIFS({V_TOTAL_RANGE},{V_CLIENTE_RANGE},$B{rr})*Configuracion!$C${ROW_FACTOR}')
    for c in (2, 3):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c == 3:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
CLI_SUMMARY_LAST = CLI_SUMMARY_FIRST + (CLI_LAST - CLI_FIRST)

# ---- Ventas mensuales ----
r2 = CLI_SUMMARY_LAST + 3
ws.cell(row=r2, column=2, value="VENTAS MENSUALES")
ws.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=4)
ws.cell(row=r2, column=2).font = f_h2
for c in range(2, 5):
    ws.cell(row=r2, column=c).fill = fill_header2
ws.row_dimensions[r2].height = 20

hr3 = r2 + 1
ws.cell(row=hr3, column=2, value="Mes")
ws.cell(row=hr3, column=3, value="Total ventas")
ws.cell(row=hr3, column=4, value="N° de ventas")
style_header_row(ws, hr3, 2, 4, height=18)

MONTH_FIRST = hr3 + 1
for i, mname in enumerate(MONTHS_ES):
    rr = MONTH_FIRST + i
    mnum = i + 1
    start_expr = f'DATE(Configuracion!$C${ROW_ANIO},{mnum},1)'
    end_expr = f'EOMONTH(DATE(Configuracion!$C${ROW_ANIO},{mnum},1),0)'
    ws.cell(row=rr, column=2, value=mname)
    ws.cell(row=rr, column=3,
            value=(f'=SUMIFS({V_TOTAL_RANGE},{V_FECHA_RANGE},">="&{start_expr},{V_FECHA_RANGE},"<="&{end_expr})'
                   f'*Configuracion!$C${ROW_FACTOR}'))
    ws.cell(row=rr, column=4,
            value=f'=COUNTIFS({V_FECHA_RANGE},">="&{start_expr},{V_FECHA_RANGE},"<="&{end_expr})')
    for c in (2, 3, 4):
        cell = ws.cell(row=rr, column=c)
        cell.border = border_all
        cell.font = f_body
        if c == 3:
            cell.number_format = "#,##0.00"
            cell.alignment = align_center
        elif c == 4:
            cell.alignment = align_center
MONTH_LAST = MONTH_FIRST + len(MONTHS_ES) - 1

tot_row = MONTH_LAST + 1
ws.cell(row=tot_row, column=2, value="TOTAL").font = f_body_b
for c, col_letter in zip((3, 4), ("C", "D")):
    cell = ws.cell(row=tot_row, column=c, value=f"=SUM({col_letter}{MONTH_FIRST}:{col_letter}{MONTH_LAST})")
    cell.number_format = "#,##0.00" if c == 3 else "#,##0"
    cell.font = f_body_b
    cell.fill = fill_light
    cell.alignment = align_center
for c in range(2, 5):
    ws.cell(row=tot_row, column=c).border = border_all
    ws.cell(row=tot_row, column=c).fill = fill_light
VENTAS_TOTAL_ROW = tot_row

# ---- Extra KPI helpers ----
kpi_helper_row = tot_row + 2
ws.cell(row=kpi_helper_row, column=2, value="Total pendiente de cobro")
ws.cell(row=kpi_helper_row, column=3, value=f'=SUMIFS({V_TOTAL_RANGE},{V_ESTADO_RANGE},"Pendiente")*Configuracion!$C${ROW_FACTOR}')
ws.cell(row=kpi_helper_row, column=3).number_format = "#,##0.00"
ws.cell(row=kpi_helper_row + 1, column=2, value="Ticket promedio")
ws.cell(row=kpi_helper_row + 1, column=3,
        value=f'=IF(D{VENTAS_TOTAL_ROW}=0,0,C{VENTAS_TOTAL_ROW}/D{VENTAS_TOTAL_ROW})')
ws.cell(row=kpi_helper_row + 1, column=3).number_format = "#,##0.00"
ws.cell(row=kpi_helper_row + 2, column=2, value="Producto más vendido (por monto)")
ws.cell(row=kpi_helper_row + 2, column=3,
        value=f'=INDEX(B{PROD_SUMMARY_FIRST}:B{PROD_SUMMARY_LAST},MATCH(MAX(D{PROD_SUMMARY_FIRST}:D{PROD_SUMMARY_LAST}),D{PROD_SUMMARY_FIRST}:D{PROD_SUMMARY_LAST},0))')
for rr in (kpi_helper_row, kpi_helper_row + 1, kpi_helper_row + 2):
    ws.cell(row=rr, column=2).font = f_body_b
    ws.cell(row=rr, column=2).border = border_all
    ws.cell(row=rr, column=3).border = border_all
    ws.cell(row=rr, column=3).font = f_body
    ws.cell(row=rr, column=3).alignment = align_center
ROW_PENDIENTE = kpi_helper_row
ROW_TICKET = kpi_helper_row + 1
ROW_TOP_PRODUCTO = kpi_helper_row + 2

lock_all(ws)

# =========================================================
# DASHBOARD
# =========================================================
ws = wb.create_sheet("Dashboard")
set_col_widths(ws, [3, 16, 16, 16, 16, 16, 16, 16, 3])
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:H2")
ws["B2"] = "DASHBOARD · CONTROL DE VENTAS"
ws["B2"].font = f_h1
ws.row_dimensions[2].height = 28
ws.merge_cells("B3:H3")
ws["B3"] = '=" Moneda: " & Configuracion!$C$8 & "   |   Año: " & Configuracion!$C$' + str(ROW_ANIO)
ws["B3"].font = f_note

kpi_row = 5
card_defs = [
    ("VENTAS DEL AÑO", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!C{VENTAS_TOTAL_ROW},\"#,##0.00\")", TEAL),
    ("TICKET PROMEDIO", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!C{ROW_TICKET},\"#,##0.00\")", NAVY),
    ("PENDIENTE DE COBRO", f"=Configuracion!$C${ROW_SIMBOLO_ACTUAL}&\" \"&TEXT(Resumen!C{ROW_PENDIENTE},\"#,##0.00\")", RED_NEG),
    ("PRODUCTO TOP", f"=Resumen!C{ROW_TOP_PRODUCTO}", GOLD),
]
col_starts = [2, 4, 6, 8]
for (label, formula, color), col0 in zip(card_defs, col_starts):
    ws.merge_cells(start_row=kpi_row, start_column=col0, end_row=kpi_row, end_column=col0 + 1)
    lc = ws.cell(row=kpi_row, column=col0, value=label)
    lc.font = f_kpi_label
    lc.alignment = align_center
    for c in range(col0, col0 + 2):
        ws.cell(row=kpi_row, column=c).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=kpi_row + 1, start_column=col0, end_row=kpi_row + 2, end_column=col0 + 1)
    vc = ws.cell(row=kpi_row + 1, column=col0, value=formula)
    if label == "PRODUCTO TOP":
        vc.font = Font(name=FONT_NAME, size=11, bold=True, color=WHITE)
        vc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        vc.font = f_kpi_value
        vc.alignment = align_center
    for rr in (kpi_row + 1, kpi_row + 2):
        for c in range(col0, col0 + 2):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor=color)

ws.row_dimensions[kpi_row].height = 18
ws.row_dimensions[kpi_row + 1].height = 22
ws.row_dimensions[kpi_row + 2].height = 22

chart_top = kpi_row + 4

pie = PieChart()
pie.title = "Ventas por producto"
data_pie = Reference(wb["Resumen"], min_col=4, min_row=hr, max_row=PROD_SUMMARY_LAST)
cats_pie = Reference(wb["Resumen"], min_col=2, min_row=PROD_SUMMARY_FIRST, max_row=PROD_SUMMARY_LAST)
pie.add_data(data_pie, titles_from_data=True)
pie.set_categories(cats_pie)
pie.height = 9
pie.width = 15
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showVal = False
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showLegendKey = False
pie.dataLabels.showBubbleSize = False
ws.add_chart(pie, f"B{chart_top}")

bar = BarChart()
bar.type = "col"
bar.title = "Ventas por mes"
bar.style = 10
bdata = Reference(wb["Resumen"], min_col=3, min_row=hr3, max_row=MONTH_LAST)
bcats = Reference(wb["Resumen"], min_col=2, min_row=MONTH_FIRST, max_row=MONTH_LAST)
bar.add_data(bdata, titles_from_data=True)
bar.set_categories(bcats)
bar.height = 9
bar.width = 15
ws.add_chart(bar, f"E{chart_top}")

bar2 = BarChart()
bar2.type = "bar"
bar2.title = "Top clientes por compras"
bar2.style = 12
bdata2 = Reference(wb["Resumen"], min_col=3, min_row=hr2, max_row=CLI_SUMMARY_LAST)
bcats2 = Reference(wb["Resumen"], min_col=2, min_row=CLI_SUMMARY_FIRST, max_row=CLI_SUMMARY_LAST)
bar2.add_data(bdata2, titles_from_data=True)
bar2.set_categories(bcats2)
bar2.height = 9
bar2.width = 30
ws.add_chart(bar2, f"B{chart_top + 19}")

lock_all(ws)

# =========================================================
wb.defined_names["TipoCambio"] = DefinedName("TipoCambio", attr_text=f"Configuracion!$C${cfg['ROW_TC']}")
wb.defined_names["MonedaVisualizacion"] = DefinedName("MonedaVisualizacion", attr_text=f"Configuracion!$C${ROW_VIS}")
wb.defined_names["FactorConversion"] = DefinedName("FactorConversion", attr_text=f"Configuracion!$C${ROW_FACTOR}")

order = ["Portada", "Instrucciones", "Configuracion", "Catalogo", "Clientes", "Ventas", "Resumen", "Dashboard"]
wb._sheets = [wb[name] for name in order]
tab_colors = {
    "Portada": NAVY, "Instrucciones": TEAL, "Configuracion": GOLD, "Catalogo": GOLD,
    "Clientes": GOLD, "Ventas": TEAL, "Resumen": NAVY, "Dashboard": NAVY,
}
for name, color in tab_colors.items():
    wb[name].sheet_properties.tabColor = color

wb.active = wb.index(wb["Dashboard"])
wb.calculation.fullCalcOnLoad = True

for name in order:
    psheet = wb[name]
    psheet.page_setup.orientation = "landscape"
    psheet.page_setup.fitToWidth = 1
    psheet.page_setup.fitToHeight = 0
    psheet.sheet_properties.pageSetUpPr.fitToPage = True
    psheet.page_margins.left = 0.4
    psheet.page_margins.right = 0.4
    psheet.page_margins.top = 0.5
    psheet.page_margins.bottom = 0.5

out_path = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "productos", "07-control-ventas", "Control-de-Ventas.xlsx",
)
os.makedirs(os.path.dirname(out_path), exist_ok=True)
wb.save(out_path)
print("Saved:", out_path)
print("Key rows:", dict(VENTAS_TOTAL_ROW=VENTAS_TOTAL_ROW, ROW_PENDIENTE=ROW_PENDIENTE))
